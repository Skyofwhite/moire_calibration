#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
run_hybrid_variants_minimal_checkpointed.py

- Removes pandas dependency (pure Python summarization).
- Adds checkpoint CSV (append-per-pair) + resume/finalize-only to avoid losing long runs.
- Writes final XLSX with raw rows + summary sheets.

Usage (first run):
  python run_hybrid_variants_minimal.py --root "...\\test" --variants "center_loftr,patch_loftr" --out_xlsx hybrid.xlsx

Resume after crash/interrupt:
  python run_hybrid_variants_minimal.py --root "...\\test" --variants "center_loftr,patch_loftr" --out_xlsx hybrid.xlsx --resume

Finalize only (CSV -> XLSX + summaries):
  python run_hybrid_variants_minimal.py --root "...\\test" --out_xlsx hybrid.xlsx --finalize_only
"""
import argparse
import csv
import json
import os
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Any

import numpy as np
from tqdm import tqdm
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import moirepose_calib as MP
import run_moirepose_hybrid_refine_gpu_with_progress_best as HR

PAIR_RE = re.compile(
    r"^(?P<idx>\d{4})_(?P<base>gt|moire)_(?P<augtype>blackbox|rotation|translation)_(?P<k>\d+)\.jpg$",
    re.IGNORECASE,
)
ORIG_RE = re.compile(r"^(?P<idx>\d{4})_(?P<base>gt|moire)\.jpg$", re.IGNORECASE)

# Fixed header for checkpoint CSV + XLSX rows
ROW_HEADERS = [
    "variant", "idx", "base", "augtype", "aug_k",
    "orig_path", "aug_path",

    "error",  # empty string means success
    "reason",

    "nmatch", "ninliers",
    "a0_m", "a_ref_m",

    "hybrid_rms_before_px", "hybrid_rms_after_px",

    "ref_aug_x_m", "ref_aug_y_m", "ref_aug_z_m",
    "mp_aug_internal_rms_px",

    "t_moirepose", "t_match", "t_h", "t_refine", "t_total",
]

def autosize_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 80)

def excel_safe(v: Any):
    """Convert to an openpyxl-friendly scalar (or string)."""
    if v is None:
        return None
    # numpy scalar -> python scalar
    if isinstance(v, (np.integer, np.floating, np.bool_)):
        return v.item()
    if isinstance(v, Path):
        return str(v)
    # tuple/list/ndarray/dict -> JSON string
    if isinstance(v, (tuple, list, np.ndarray)):
        return json.dumps([excel_safe(x) for x in list(v)], ensure_ascii=False)
    if isinstance(v, dict):
        return json.dumps({str(k): excel_safe(val) for k, val in v.items()}, ensure_ascii=False)
    return v

def csv_safe(v: Any) -> str:
    """Convert value to a CSV-safe string (no newlines)."""
    if v is None:
        return ""
    v2 = excel_safe(v)
    s = str(v2)
    # remove newlines (keep CSV one-line per row)
    s = s.replace("\r", " ").replace("\n", " ")
    return s

def find_originals(test_dir: Path) -> Dict[Tuple[str, str], Path]:
    mapping: Dict[Tuple[str, str], Path] = {}
    for p in test_dir.iterdir():
        if p.is_file():
            m = ORIG_RE.match(p.name)
            if m:
                mapping[(m.group("idx"), m.group("base").lower())] = p
    return mapping

def find_augmented(aug_dir: Path) -> List[Tuple[str, str, str, str, Path]]:
    items: List[Tuple[str, str, str, str, Path]] = []
    if not aug_dir.exists():
        return items
    for p in aug_dir.iterdir():
        if p.is_file():
            m = PAIR_RE.match(p.name)
            if m:
                items.append((m.group("idx"), m.group("base").lower(),
                              m.group("augtype").lower(), m.group("k"), p))
    items.sort(key=lambda t: (t[0], t[1], t[2], int(t[3])))
    return items

def row_key_from_fields(variant: str, idx: str, base: str, augtype: str, aug_k: int) -> str:
    return f"{variant}|{idx}|{base}|{augtype}|{aug_k}"

def load_done_keys_from_checkpoint(csv_path: Path) -> set:
    done = set()
    if not csv_path.exists():
        return done
    with csv_path.open("r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            try:
                k = row_key_from_fields(
                    r.get("variant", ""),
                    r.get("idx", ""),
                    r.get("base", ""),
                    r.get("augtype", ""),
                    int(r.get("aug_k", "0") or 0),
                )
                done.add(k)
            except Exception:
                continue
    return done

def parse_float(x: str) -> Optional[float]:
    if x is None:
        return None
    s = str(x).strip()
    if s == "" or s.lower() == "nan":
        return None
    try:
        return float(s)
    except Exception:
        return None

def mean(vals: List[float]) -> Optional[float]:
    if not vals:
        return None
    return float(sum(vals) / len(vals))

def median(vals: List[float]) -> Optional[float]:
    if not vals:
        return None
    s = sorted(vals)
    n = len(s)
    mid = n // 2
    if n % 2 == 1:
        return float(s[mid])
    return float(0.5 * (s[mid - 1] + s[mid]))

def summarize_from_rows(rows: List[dict], group_keys: List[str]) -> List[dict]:
    """
    Pure Python grouping summary.
    Success condition: error == "" (empty) or None.
    """
    agg = {}
    for r in rows:
        gk = tuple(r.get(k) for k in group_keys)
        if gk not in agg:
            agg[gk] = {
                "N": 0,
                "success": 0,
                "before": [],
                "after": [],
                "t_total": [],
            }
        a = agg[gk]
        a["N"] += 1
        err = r.get("error")
        ok = (err is None) or (str(err).strip() == "")
        if ok:
            a["success"] += 1
            vb = r.get("hybrid_rms_before_px")
            va = r.get("hybrid_rms_after_px")
            vt = r.get("t_total")
            if vb is not None:
                a["before"].append(float(vb))
            if va is not None:
                a["after"].append(float(va))
            if vt is not None:
                a["t_total"].append(float(vt))

    out = []
    for gk, a in agg.items():
        row = {k: v for k, v in zip(group_keys, gk)}
        N = int(a["N"])
        succ = int(a["success"])
        row.update({
            "N": N,
            "success": succ,
            "success_rate": (succ / N) if N > 0 else None,
            "hybrid_rms_before_med": median(a["before"]) if succ > 0 else None,
            "hybrid_rms_after_med": median(a["after"]) if succ > 0 else None,
            "hybrid_rms_before_mean": mean(a["before"]) if succ > 0 else None,
            "hybrid_rms_after_mean": mean(a["after"]) if succ > 0 else None,
            "t_total_mean": mean(a["t_total"]) if succ > 0 else None,
        })
        out.append(row)

    def _sort_key(r):
        return tuple("" if r.get(k) is None else str(r.get(k)) for k in group_keys)
    out.sort(key=_sort_key)
    return out

def read_checkpoint_rows(csv_path: Path) -> List[dict]:
    rows = []
    if not csv_path.exists():
        return rows
    with csv_path.open("r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            rr = dict(r)
            rr["aug_k"] = int(rr.get("aug_k", "0") or 0)
            for k in ["nmatch", "ninliers", "a0_m", "a_ref_m",
                      "hybrid_rms_before_px", "hybrid_rms_after_px",
                      "ref_aug_x_m", "ref_aug_y_m", "ref_aug_z_m",
                      "mp_aug_internal_rms_px",
                      "t_moirepose", "t_match", "t_h", "t_refine", "t_total"]:
                rr[k] = parse_float(rr.get(k, ""))
            rr["error"] = (rr.get("error") or "").strip()
            rr["reason"] = (rr.get("reason") or "").strip()
            rows.append(rr)
    return rows

def write_xlsx_from_rows(out_xlsx: Path, rows: List[dict]):
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "rows"

    headers = ROW_HEADERS
    ws.append(headers)
    for r in rows:
        ws.append([excel_safe(r.get(h, None)) for h in headers])
    autosize_columns(ws)

    summary1 = summarize_from_rows(rows, ["variant"])
    summary2 = summarize_from_rows(rows, ["variant", "base", "augtype"])

    ws1 = wb.create_sheet("summary_variant")
    headers1 = sorted({k for r in summary1 for k in r.keys()}) if summary1 else ["variant","N","success","success_rate"]
    ws1.append(headers1)
    for r in summary1:
        ws1.append([excel_safe(r.get(h, None)) for h in headers1])
    autosize_columns(ws1)

    ws2 = wb.create_sheet("summary_variant_base_aug")
    headers2 = sorted({k for r in summary2 for k in r.keys()}) if summary2 else ["variant","base","augtype","N","success","success_rate"]
    ws2.append(headers2)
    for r in summary2:
        ws2.append([excel_safe(r.get(h, None)) for h in headers2])
    autosize_columns(ws2)

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_xlsx))
    print(f"[DONE] Saved XLSX: {out_xlsx} rows={len(rows)}")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--aug_dir", default="augmented_9x")
    ap.add_argument("--out_xlsx", default="hybrid_variants.xlsx")
    ap.add_argument("--checkpoint_csv", default=None, help="default: <out_xlsx>.checkpoint.csv")

    ap.add_argument("--variants", default="center_loftr,patch_loftr,center_orb,patch_orb")
    ap.add_argument("--resume", action="store_true")
    ap.add_argument("--finalize_only", action="store_true")
    ap.add_argument("--flush_every", type=int, default=1, help="flush+fsync every N rows (1=most safe)")

    # LoFTR
    ap.add_argument("--loftr_device", default="cuda")
    ap.add_argument("--loftr_pretrained", default="outdoor", choices=["outdoor", "indoor"])
    ap.add_argument("--loftr_max_side", type=int, default=640)
    ap.add_argument("--loftr_conf_th", type=float, default=0.3)

    # Homography / inliers
    ap.add_argument("--use_usac_magsac", action="store_true")
    ap.add_argument("--h_thr_px", type=float, default=3.0)
    ap.add_argument("--sym_err_th", type=float, default=6.0)
    ap.add_argument("--sym_err_keep", type=int, default=400)
    ap.add_argument("--min_inliers", type=int, default=80)
    ap.add_argument("--max_refine_points", type=int, default=300)

    # Refine
    ap.add_argument("--refine_backend", default="opencv", choices=["opencv", "torch"])
    ap.add_argument("--refine_device", default="cuda")
    ap.add_argument("--torch_iters", type=int, default=10)
    ap.add_argument("--torch_method", default="adam", choices=["adam", "lbfgs"])
    ap.add_argument("--huber_delta", type=float, default=3.0)
    ap.add_argument("--optimize_scale_a", action="store_true")

    # Intrinsics
    ap.add_argument("--use_full_k", action="store_true")
    ap.add_argument("--K_full_fx_fy_cx_cy", nargs=4, type=float, default=None)

    # Gating
    ap.add_argument("--gate_rms_before", type=float, default=60.0)
    ap.add_argument("--gate_abs_xyz_m", type=float, default=5.0)

    args = ap.parse_args()

    root = Path(args.root)
    aug_dir = root / args.aug_dir
    out_xlsx = Path(args.out_xlsx)
    if not out_xlsx.is_absolute():
        out_xlsx = root / out_xlsx

    checkpoint_csv = Path(args.checkpoint_csv) if args.checkpoint_csv else Path(str(out_xlsx) + ".checkpoint.csv")

    if args.finalize_only:
        rows = read_checkpoint_rows(checkpoint_csv)
        if not rows:
            print(f"[ERROR] No checkpoint CSV found or empty: {checkpoint_csv}")
            sys.exit(2)
        write_xlsx_from_rows(out_xlsx, rows)
        return

    originals = find_originals(root)
    augmented = find_augmented(aug_dir)

    variants = [v.strip() for v in args.variants.split(",") if v.strip()]
    need_loftr = any(v.endswith("_loftr") for v in variants)

    loftr = None
    if need_loftr:
        if not getattr(HR, "_HAS_LOFTR", False):
            raise RuntimeError("LoFTR requested but torch+kornia not available.")
        loftr = HR.LoFTRMatcher(device=args.loftr_device, pretrained=args.loftr_pretrained)

    full_fx_fy_cx_cy = None
    if args.use_full_k:
        if args.K_full_fx_fy_cx_cy is None:
            raise RuntimeError("--use_full_k requires --K_full_fx_fy_cx_cy fx fy cx cy")
        full_fx_fy_cx_cy = tuple(map(float, args.K_full_fx_fy_cx_cy))

    done_keys = set()
    if args.resume:
        done_keys = load_done_keys_from_checkpoint(checkpoint_csv)
        print(f"[INFO] resume enabled. done_keys={len(done_keys)} from {checkpoint_csv}")

    need_header = not checkpoint_csv.exists()
    checkpoint_csv.parent.mkdir(parents=True, exist_ok=True)
    f = checkpoint_csv.open("a", newline="", encoding="utf-8")
    writer = csv.DictWriter(f, fieldnames=ROW_HEADERS)
    if need_header:
        writer.writeheader()
        f.flush()
        os.fsync(f.fileno())

    print(f"[INFO] root={root} pairs={len(augmented)} variants={variants}")
    print(f"[INFO] checkpoint_csv={checkpoint_csv}")

    rows_written = 0
    try:
        for variant in variants:
            if variant == "center_loftr":
                patch_mode, matcher = False, "loftr"
            elif variant == "patch_loftr":
                patch_mode, matcher = True, "loftr"
            elif variant == "center_orb":
                patch_mode, matcher = False, "orb"
            elif variant == "patch_orb":
                patch_mode, matcher = True, "orb"
            else:
                raise RuntimeError(f"Unknown variant: {variant}")

            for (idx, base, augtype, k, aug_path) in tqdm(augmented, desc=variant, unit="pair"):
                aug_k = int(k)
                key = row_key_from_fields(variant, idx, base, augtype, aug_k)
                if key in done_keys:
                    continue

                orig_path = originals.get((idx, base))
                if orig_path is None:
                    row = {h: "" for h in ROW_HEADERS}
                    row.update({
                        "variant": variant, "idx": idx, "base": base, "augtype": augtype, "aug_k": aug_k,
                        "orig_path": "", "aug_path": str(aug_path),
                        "error": "missing_original",
                        "reason": f"missing original {idx}_{base}.jpg",
                    })
                    writer.writerow({h: csv_safe(row.get(h, "")) for h in ROW_HEADERS})
                    rows_written += 1
                    done_keys.add(key)
                    if rows_written % max(args.flush_every, 1) == 0:
                        f.flush()
                        os.fsync(f.fileno())
                    continue

                out = HR.run_one_pair(
                    orig_path=orig_path,
                    aug_path=aug_path,
                    params=MP.DEVICE_PARAMS,
                    patch_mode=patch_mode,
                    matcher=matcher,
                    loftr=loftr,
                    loftr_max_side=args.loftr_max_side,
                    loftr_conf_th=args.loftr_conf_th,
                    use_usac_magsac=args.use_usac_magsac,
                    h_thr_px=args.h_thr_px,
                    sym_err_th=args.sym_err_th,
                    sym_err_keep=args.sym_err_keep,
                    min_inliers=args.min_inliers,
                    max_refine_points=args.max_refine_points,
                    refine_backend=args.refine_backend,
                    refine_device=args.refine_device,
                    torch_iters=args.torch_iters,
                    torch_method=args.torch_method,
                    huber_delta=args.huber_delta,
                    optimize_scale_a=args.optimize_scale_a,
                    use_full_k=args.use_full_k,
                    full_fx_fy_cx_cy=full_fx_fy_cx_cy,
                    gate_rms_before=args.gate_rms_before,
                    gate_abs_xyz_m=args.gate_abs_xyz_m,
                )

                row = {h: "" for h in ROW_HEADERS}
                row.update({
                    "variant": variant,
                    "idx": idx, "base": base, "augtype": augtype, "aug_k": aug_k,
                    "orig_path": str(orig_path), "aug_path": str(aug_path),
                    "nmatch": out.get("nmatch"),
                    "ninliers": out.get("ninliers"),
                    "hybrid_rms_before_px": out.get("rms_before_px"),
                    "hybrid_rms_after_px": out.get("rms_after_px"),
                    "t_moirepose": out.get("t_moirepose"),
                    "t_match": out.get("t_match"),
                    "t_h": out.get("t_h"),
                    "t_refine": out.get("t_refine"),
                    "t_total": out.get("t_total"),
                    "a0_m": out.get("a0_m"),
                    "a_ref_m": out.get("a_ref_m"),
                })

                if not out.get("ok", False):
                    row["error"] = "fail"
                    row["reason"] = out.get("reason", "")
                else:
                    try:
                        pair = out.get("pair", {})
                        aug_res = pair.get("pair_result", {}).get("aug", {})
                        row["mp_aug_internal_rms_px"] = aug_res.get("rms_reproj_px", None)
                    except Exception:
                        pass
                    rp = out.get("refined_pose", {})
                    row["ref_aug_x_m"] = rp.get("x")
                    row["ref_aug_y_m"] = rp.get("y")
                    row["ref_aug_z_m"] = rp.get("z")
                    row["error"] = ""  # success

                writer.writerow({h: csv_safe(row.get(h, "")) for h in ROW_HEADERS})
                rows_written += 1
                done_keys.add(key)

                if rows_written % max(args.flush_every, 1) == 0:
                    f.flush()
                    os.fsync(f.fileno())

    except KeyboardInterrupt:
        print("\n[WARN] Interrupted by user (Ctrl+C). Checkpoint is safe. You can resume with --resume.")
    finally:
        try:
            f.flush()
            os.fsync(f.fileno())
        except Exception:
            pass
        f.close()

    rows = read_checkpoint_rows(checkpoint_csv)
    if not rows:
        print(f"[ERROR] No rows in checkpoint CSV: {checkpoint_csv}")
        sys.exit(2)
    write_xlsx_from_rows(out_xlsx, rows)

if __name__ == "__main__":
    main()
