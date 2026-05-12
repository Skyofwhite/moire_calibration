import argparse
import re
import time
from pathlib import Path
from typing import Dict, List, Tuple, Optional

import cv2
import numpy as np
from tqdm import tqdm
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import moirepose_calib as MP  # your file

# Optional LoFTR (Kornia)
try:
    import torch
    import kornia.feature as KF
    _HAS_TORCH = True
    _HAS_LOFTR = True
except Exception:
    torch = None
    KF = None
    _HAS_TORCH = False
    _HAS_LOFTR = False

def _inference_mode_decorator():
    if _HAS_TORCH and torch is not None:
        return torch.inference_mode()
    def _decorator(fn):
        return fn
    return _decorator

PAIR_RE = re.compile(
    r"^(?P<idx>\d{4})_(?P<base>gt|moire)_(?P<augtype>blackbox|rotation|translation)_(?P<k>\d+)\.jpg$",
    re.IGNORECASE
)
ORIG_RE = re.compile(r"^(?P<idx>\d{4})_(?P<base>gt|moire)\.jpg$", re.IGNORECASE)


# -----------------------------
# Excel helper
# -----------------------------
def autosize_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


# -----------------------------
# Pair discovery
# -----------------------------
def find_originals(test_dir: Path) -> Dict[Tuple[str, str], Path]:
    mapping: Dict[Tuple[str, str], Path] = {}
    for p in test_dir.iterdir():
        if p.is_file():
            m = ORIG_RE.match(p.name)
            if m:
                mapping[(m.group("idx"), m.group("base").lower())] = p
    return mapping


def find_augmented(aug_dir: Path) -> List[Tuple[str, str, str, str, Path]]:
    items = []
    if not aug_dir.exists():
        return items
    for p in aug_dir.iterdir():
        if p.is_file():
            m = PAIR_RE.match(p.name)
            if m:
                items.append((m.group("idx"), m.group("base").lower(),
                              m.group("augtype").lower(), m.group("k"), p))
    items.sort(key=lambda t: (t[0], t[1], t[2], int(t[3])))
    return items


# -----------------------------
# IO and crop
# -----------------------------
def load_bgr(path: Path) -> np.ndarray:
    img = cv2.imread(str(path), cv2.IMREAD_COLOR)
    if img is None:
        raise FileNotFoundError(str(path))
    return img


def crop_gray(img_bgr: np.ndarray, box: Tuple[int, int, int, int]) -> np.ndarray:
    x0, y0, x1, y1 = box
    crop = img_bgr[y0:y1, x0:x1]
    if crop.size == 0:
        raise ValueError(f"Empty crop for box={box} on image size={img_bgr.shape}")
    return cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)


# -----------------------------
# ORB matching (CPU)
# -----------------------------
def orb_match_points(
    g0: np.ndarray,
    g1: np.ndarray,
    nfeatures: int = 3000,
    max_matches: int = 1500
):
    orb = cv2.ORB_create(nfeatures=nfeatures)
    kp0, des0 = orb.detectAndCompute(g0, None)
    kp1, des1 = orb.detectAndCompute(g1, None)

    nkp0, nkp1 = len(kp0), len(kp1)
    if des0 is None or des1 is None or nkp0 < 30 or nkp1 < 30:
        return (np.zeros((0, 2), np.float64),
                np.zeros((0, 2), np.float64),
                nkp0, nkp1, 0)

    bf = cv2.BFMatcher(cv2.NORM_HAMMING, crossCheck=True)
    matches = bf.match(des0, des1)
    matches = sorted(matches, key=lambda m: m.distance)[:max_matches]

    pts0 = np.float64([kp0[m.queryIdx].pt for m in matches]) if matches else np.zeros((0, 2), np.float64)
    pts1 = np.float64([kp1[m.trainIdx].pt for m in matches]) if matches else np.zeros((0, 2), np.float64)
    return pts0, pts1, nkp0, nkp1, len(matches)


# -----------------------------
# LoFTR matching (GPU via Kornia)
# -----------------------------
class LoFTRMatcher:
    def __init__(self, device: str = "cuda", pretrained: str = "outdoor"):
        if not _HAS_LOFTR:
            raise RuntimeError("LoFTR requested but kornia/torch is not installed.")
        if device.startswith("cuda") and (not torch.cuda.is_available()):
            raise RuntimeError("CUDA requested but torch.cuda.is_available() is False.")
        self.device = torch.device(device)
        self.matcher = KF.LoFTR(pretrained=pretrained).to(self.device).eval()

    @_inference_mode_decorator()
    def match(self, g0: np.ndarray, g1: np.ndarray, max_side: int = 640, conf_th: float = 0.0):
        """
        Returns mkpts0, mkpts1, nmatch, s0, s1, mean_conf
        Conf filter is applied if kornia provides a confidence tensor.
        """
        def resize_keep(img, max_side_):
            h, w = img.shape[:2]
            s = max_side_ / max(h, w)
            if s >= 1.0:
                return img, 1.0
            out = cv2.resize(img, (int(round(w * s)), int(round(h * s))), interpolation=cv2.INTER_AREA)
            return out, s

        g0r, s0 = resize_keep(g0, max_side)
        g1r, s1 = resize_keep(g1, max_side)

        t0 = torch.from_numpy(g0r).float()[None, None] / 255.0
        t1 = torch.from_numpy(g1r).float()[None, None] / 255.0
        t0 = t0.to(self.device)
        t1 = t1.to(self.device)

        out = self.matcher({"image0": t0, "image1": t1})

        mkpts0 = out["keypoints0"].detach().cpu().numpy()
        mkpts1 = out["keypoints1"].detach().cpu().numpy()

        conf = None
        for k in ["confidence", "conf", "scores"]:
            if k in out:
                conf = out[k].detach().cpu().numpy().reshape(-1)
                break

        mean_conf = float(np.mean(conf)) if conf is not None and conf.size > 0 else None

        if conf is not None and conf_th > 0.0:
            m = conf >= conf_th
            mkpts0 = mkpts0[m]
            mkpts1 = mkpts1[m]
            conf = conf[m]

        # back to original crop coords
        if mkpts0.shape[0] > 0:
            mkpts0[:, 0] /= s0
            mkpts0[:, 1] /= s0
            mkpts1[:, 0] /= s1
            mkpts1[:, 1] /= s1

        return mkpts0.astype(np.float64), mkpts1.astype(np.float64), int(mkpts0.shape[0]), float(s0), float(s1), mean_conf


# -----------------------------
# Homography + symmetric transfer error filtering
# -----------------------------
def _apply_H(H: np.ndarray, pts: np.ndarray) -> np.ndarray:
    pts_h = np.concatenate([pts, np.ones((pts.shape[0], 1), dtype=np.float64)], axis=1)  # (N,3)
    q = (H @ pts_h.T).T
    q = q[:, :2] / np.clip(q[:, 2:3], 1e-12, None)
    return q


def symmetric_transfer_errors(H: np.ndarray, pts0: np.ndarray, pts1: np.ndarray) -> np.ndarray:
    """
    returns per-point symmetric transfer error in pixels:
      e = ||H p0 - p1|| + ||H^{-1} p1 - p0||
    """
    if H is None:
        return np.full((pts0.shape[0],), np.inf, dtype=np.float64)
    try:
        Hinv = np.linalg.inv(H)
    except np.linalg.LinAlgError:
        return np.full((pts0.shape[0],), np.inf, dtype=np.float64)

    p01 = _apply_H(H, pts0)
    p10 = _apply_H(Hinv, pts1)

    e01 = np.linalg.norm(p01 - pts1, axis=1)
    e10 = np.linalg.norm(p10 - pts0, axis=1)
    return e01 + e10


def robust_h_inliers(
    pts0: np.ndarray,
    pts1: np.ndarray,
    thr_px: float,
    use_usac_magsac: bool,
    sym_err_th: float,
    sym_err_keep: int
):
    if pts0.shape[0] < 50:
        return np.zeros((0, 2), np.float64), np.zeros((0, 2), np.float64), 0, None, None

    method = cv2.USAC_MAGSAC if (use_usac_magsac and hasattr(cv2, "USAC_MAGSAC")) else cv2.RANSAC
    H, mask = cv2.findHomography(pts0, pts1, method=method, ransacReprojThreshold=thr_px)
    if H is None or mask is None:
        return np.zeros((0, 2), np.float64), np.zeros((0, 2), np.float64), 0, None, None

    m = mask.ravel().astype(bool)
    in0 = pts0[m]
    in1 = pts1[m]
    if in0.shape[0] == 0:
        return np.zeros((0, 2), np.float64), np.zeros((0, 2), np.float64), 0, H, None

    # symmetric error filter (2nd stage)
    sym = symmetric_transfer_errors(H, in0, in1)
    if np.isfinite(sym).sum() == 0:
        return np.zeros((0, 2), np.float64), np.zeros((0, 2), np.float64), 0, H, None

    if sym_err_th > 0:
        keep = sym <= sym_err_th
        in0b, in1b, symb = in0[keep], in1[keep], sym[keep]
    else:
        in0b, in1b, symb = in0, in1, sym

    # keep top-K smallest symmetric errors (stabilizes repeated pattern outliers)
    if sym_err_keep > 0 and in0b.shape[0] > sym_err_keep:
        idx = np.argsort(symb)[:sym_err_keep]
        in0b, in1b, symb = in0b[idx], in1b[idx], symb[idx]

    return in0b, in1b, int(in0b.shape[0]), H, float(np.mean(symb)) if symb.size > 0 else None


# -----------------------------
# Full K -> crop K
# -----------------------------
def crop_intrinsics(
    crop_wh: Tuple[int, int],
    crop_box: Tuple[int, int, int, int],
    use_full_k: bool,
    full_fx_fy_cx_cy: Optional[Tuple[float, float, float, float]],
    params: dict
) -> np.ndarray:
    w, h = crop_wh
    x0, y0, x1, y1 = crop_box

    if use_full_k:
        if full_fx_fy_cx_cy is None:
            raise ValueError("use_full_k=True but full_fx_fy_cx_cy not provided")
        fx, fy, cx_full, cy_full = full_fx_fy_cx_cy
        cx_crop = cx_full - float(x0)
        cy_crop = cy_full - float(y0)
        return np.array([[fx, 0.0, cx_crop],
                         [0.0, fy, cy_crop],
                         [0.0, 0.0, 1.0]], dtype=np.float64)

    # fallback: MP internal
    f_px = MP.focal_px(float(params["f_m"]), float(params["cfa_pitch_m"]))
    cx, cy = w / 2.0, h / 2.0
    return np.array([[f_px, 0.0, cx],
                     [0.0, f_px, cy],
                     [0.0, 0.0, 1.0]], dtype=np.float64)


# -----------------------------
# MP pose <-> OpenCV
# -----------------------------
def mp_pose_to_opencv_rt(pose: dict) -> Tuple[np.ndarray, np.ndarray]:
    C = np.array([pose["x"], pose["y"], pose["z"]], dtype=np.float64).reshape(3, 1)
    u = np.array(pose["u"], dtype=np.float64).reshape(3)
    v = np.array(pose["v"], dtype=np.float64).reshape(3)
    w = np.array(pose["w"], dtype=np.float64).reshape(3)

    R_wc = np.stack([u, v, w], axis=1)
    R_cw = R_wc.T
    t = -(R_cw @ C)
    rvec, _ = cv2.Rodrigues(R_cw)
    return rvec.astype(np.float64), t.astype(np.float64)


def opencv_rt_to_mp_pose(rvec: np.ndarray, tvec: np.ndarray):
    R_cw, _ = cv2.Rodrigues(rvec.reshape(3, 1))
    R_wc = R_cw.T
    C = -(R_wc @ tvec.reshape(3, 1))
    u = R_wc[:, 0]
    v = R_wc[:, 1]
    w = R_wc[:, 2]
    return C.reshape(3), u.reshape(3), v.reshape(3), w.reshape(3)


# -----------------------------
# Mapping crop pixels -> plane points
# -----------------------------
def crop_px_to_screen_xy_np(px: np.ndarray, crop_wh: Tuple[int, int], a_m: float, margin_px: float = 4.0):
    """Legacy Eq. (13) fronto-parallel lift.

    Kept as an ablation/baseline because reviewers specifically questioned this
    approximation.  The default path below now uses ray-plane intersection.
    """
    w, h = crop_wh
    cx, cy = w / 2.0, h / 2.0
    half = (min(w, h) / 2.0) - float(margin_px)
    if half <= 1:
        half = min(w, h) / 2.0
    s = a_m / (half + 1e-12)
    x = (px[:, 0] - cx) * s
    y = -(px[:, 1] - cy) * s
    z = np.zeros_like(x)
    return np.stack([x, y, z], axis=1).astype(np.float64)


def ray_plane_lift_crop_pixels_np(
    px_crop: np.ndarray,
    K_crop: np.ndarray,
    rvec_screen_to_cam: np.ndarray,
    tvec_screen_to_cam: np.ndarray,
    plane_z: float = 0.0,
) -> Tuple[np.ndarray, np.ndarray]:
    """Back-project crop pixels and intersect each ray with the screen plane.

    This replaces the paper's Eq. (13) local fronto-parallel lift in the default
    refinement path.  It uses the moiré-initialized pose for the original view,
    so each 2D match in image0 becomes a metric 3D point on Z=0 in screen
    coordinates.  The returned validity mask removes rays parallel to the plane
    and intersections behind the camera.
    """
    if px_crop.size == 0:
        return np.zeros((0, 3), dtype=np.float64), np.zeros((0,), dtype=bool)

    Kinv = np.linalg.inv(K_crop.astype(np.float64))
    R_cw, _ = cv2.Rodrigues(rvec_screen_to_cam.reshape(3, 1))
    R_wc = R_cw.T
    C = -(R_wc @ tvec_screen_to_cam.reshape(3, 1)).reshape(3)

    pts_h = np.concatenate([px_crop.astype(np.float64), np.ones((px_crop.shape[0], 1))], axis=1)
    rays_cam = (Kinv @ pts_h.T).T
    rays_world = (R_wc @ rays_cam.T).T

    denom = rays_world[:, 2]
    valid = np.isfinite(denom) & (np.abs(denom) > 1e-9)
    lam = np.full((px_crop.shape[0],), np.nan, dtype=np.float64)
    lam[valid] = (float(plane_z) - C[2]) / denom[valid]
    valid = valid & np.isfinite(lam) & (lam > 0)

    X = C[None, :] + lam[:, None] * rays_world
    X[:, 2] = float(plane_z)
    X[~valid] = np.nan
    return X.astype(np.float64), valid


# -----------------------------
# RMS
# -----------------------------
def project_points_cv(obj_pts: np.ndarray, rvec: np.ndarray, tvec: np.ndarray, K: np.ndarray) -> np.ndarray:
    proj, _ = cv2.projectPoints(obj_pts, rvec, tvec, K, None)
    return proj.reshape(-1, 2)


def rms_px(obs: np.ndarray, pred: np.ndarray) -> float:
    e = obs - pred
    return float(np.sqrt(np.mean(np.sum(e * e, axis=1))))


# ============================================================
# Torch refine: pose + optional scale a
# ============================================================
def _torch_require():
    if not _HAS_TORCH:
        raise RuntimeError("Torch refine requested but torch is not installed.")
    return True


def _skew(w: "torch.Tensor") -> "torch.Tensor":
    wx, wy, wz = w[..., 0], w[..., 1], w[..., 2]
    O = torch.zeros_like(wx)
    K = torch.stack([
        torch.stack([O, -wz, wy], dim=-1),
        torch.stack([wz, O, -wx], dim=-1),
        torch.stack([-wy, wx, O], dim=-1),
    ], dim=-2)
    return K


def _rodrigues(w: "torch.Tensor") -> "torch.Tensor":
    theta = torch.linalg.norm(w) + 1e-12
    k = w / theta
    K = _skew(k)
    I = torch.eye(3, device=w.device, dtype=w.dtype)
    sin_t = torch.sin(theta)
    cos_t = torch.cos(theta)
    return I + sin_t * K + (1.0 - cos_t) * (K @ K)


def _project_torch(K: "torch.Tensor", R: "torch.Tensor", t: "torch.Tensor", X: "torch.Tensor") -> "torch.Tensor":
    Xc = (X @ R.T) + t[None, :]
    z = Xc[:, 2].clamp_min(1e-6)
    x = Xc[:, 0] / z
    y = Xc[:, 1] / z
    u = K[0, 0] * x + K[0, 2]
    v = K[1, 1] * y + K[1, 2]
    return torch.stack([u, v], dim=-1)


def _huber_loss(residual: "torch.Tensor", delta: float = 3.0) -> "torch.Tensor":
    r = torch.linalg.norm(residual, dim=-1)
    abs_r = torch.abs(r)
    d = torch.tensor(delta, device=r.device, dtype=r.dtype)
    quad = torch.minimum(abs_r, d)
    lin = abs_r - quad
    return (0.5 * quad * quad + d * lin).mean()


def _crop_px_to_screen_xy_torch(px: "torch.Tensor", crop_wh: Tuple[int, int], a: "torch.Tensor", margin_px: float = 4.0):
    w, h = crop_wh
    cx, cy = w / 2.0, h / 2.0
    half = (min(w, h) / 2.0) - float(margin_px)
    if half <= 1:
        half = min(w, h) / 2.0
    s = a / (half + 1e-12)
    x = (px[:, 0] - cx) * s
    y = -(px[:, 1] - cy) * s
    z = torch.zeros_like(x)
    return torch.stack([x, y, z], dim=-1)


def refine_pose_torch(
    pts0_px_np: np.ndarray,
    pts1_px_np: np.ndarray,
    K_np: np.ndarray,
    rvec0_np: np.ndarray,
    tvec0_np: np.ndarray,
    a0: float,
    crop_wh: Tuple[int, int],
    optimize_scale_a: bool,
    device: str,
    iters: int,
    huber_delta: float,
    method: str,  # "adam" or "lbfgs"
    obj_pts_fixed_np: Optional[np.ndarray] = None,
):
    _torch_require()
    if device.startswith("cuda") and (not torch.cuda.is_available()):
        raise RuntimeError("refine_device=cuda requested but torch.cuda.is_available() is False")

    dev = torch.device(device)
    dtype = torch.float32

    p0 = torch.from_numpy(pts0_px_np).to(dev, dtype=dtype)
    p1 = torch.from_numpy(pts1_px_np).to(dev, dtype=dtype)
    Kt = torch.from_numpy(K_np).to(dev, dtype=dtype)
    X_fixed = None
    if obj_pts_fixed_np is not None:
        X_fixed = torch.from_numpy(obj_pts_fixed_np.astype(np.float32)).to(dev, dtype=dtype)

    w = torch.tensor(rvec0_np.reshape(3), device=dev, dtype=dtype, requires_grad=True)
    t = torch.tensor(tvec0_np.reshape(3), device=dev, dtype=dtype, requires_grad=True)

    if optimize_scale_a and X_fixed is not None:
        # Scale optimization is only meaningful for the legacy fronto-parallel
        # object-point parameterization.  Ray-plane lifted points are already in
        # metric screen coordinates.
        optimize_scale_a = False

    if optimize_scale_a:
        log_a = torch.tensor(np.log(max(a0, 1e-12)), device=dev, dtype=dtype, requires_grad=True)
        params = [w, t, log_a]
    else:
        log_a = None
        params = [w, t]

    final_loss = None
    final_a = None

    def forward_loss():
        nonlocal final_loss, final_a
        R = _rodrigues(w)
        if X_fixed is not None:
            X = X_fixed
            a = torch.tensor(a0, device=dev, dtype=dtype)
        elif optimize_scale_a:
            a = torch.exp(log_a)
            X = _crop_px_to_screen_xy_torch(p0, crop_wh=crop_wh, a=a)
        else:
            a = torch.tensor(a0, device=dev, dtype=dtype)
            X = _crop_px_to_screen_xy_torch(p0, crop_wh=crop_wh, a=a)
        u_pred = _project_torch(Kt, R, t, X)
        loss = _huber_loss(u_pred - p1, delta=huber_delta)
        final_loss = loss.detach()
        final_a = a.detach()
        return loss

    if method == "adam":
        opt = torch.optim.Adam(params, lr=5e-3)
        for _ in range(iters):
            opt.zero_grad(set_to_none=True)
            loss = forward_loss()
            loss.backward()
            opt.step()
    else:
        opt = torch.optim.LBFGS(params, lr=1.0, max_iter=iters, line_search_fn="strong_wolfe")

        def closure():
            opt.zero_grad(set_to_none=True)
            loss = forward_loss()
            loss.backward()
            return loss

        opt.step(closure)
        forward_loss()

    rvec_ref = w.detach().cpu().numpy().astype("float64").reshape(3, 1)
    tvec_ref = t.detach().cpu().numpy().astype("float64").reshape(3, 1)
    a_ref = float(final_a.cpu().item()) if final_a is not None else float(a0)
    loss_val = float(final_loss.cpu().item()) if final_loss is not None else float("nan")
    return rvec_ref, tvec_ref, a_ref, loss_val


# ============================================================
# One pair pipeline
# ============================================================
def run_one_pair(
    orig_path: Path,
    aug_path: Path,
    params: dict,
    patch_mode: bool,
    matcher: str,
    loftr: Optional[LoFTRMatcher],
    loftr_max_side: int,
    loftr_conf_th: float,
    use_usac_magsac: bool,
    h_thr_px: float,
    sym_err_th: float,
    sym_err_keep: int,
    min_inliers: int,
    max_refine_points: int,
    refine_backend: str,  # "opencv" or "torch"
    refine_device: str,
    torch_iters: int,
    torch_method: str,
    huber_delta: float,
    optimize_scale_a: bool,
    use_full_k: bool,
    full_fx_fy_cx_cy: Optional[Tuple[float, float, float, float]],
    gate_rms_before: float,
    gate_abs_xyz_m: float,
    object_lift: str = "ray_plane",  # "ray_plane" or "fronto_parallel"
    patch_grid: Tuple[int, int] = (5, 5),
    patch_score_mode: str = "reliability",
) -> Dict:
    t_total0 = time.time()

    # 1) MP initial
    t0 = time.time()
    pair = MP.calibrate_pair(
        str(orig_path),
        str(aug_path),
        params,
        patch_mode=patch_mode,
        patch_grid=patch_grid,
        patch_score_mode=patch_score_mode,
    )
    t1 = time.time()

    if "pair_result" not in pair:
        return {"ok": False, "reason": "MP.calibrate_pair returned unexpected structure"}

    orig_res = pair["pair_result"].get("orig", {})
    aug_res = pair["pair_result"].get("aug", {})

    # crop box
    if patch_mode:
        crop_box = pair.get("debug", {}).get("chosen_patch_box", None)
        if crop_box is None:
            return {"ok": False, "reason": "patch_mode=True but debug.chosen_patch_box missing"}
    else:
        crop_box = orig_res.get("center_crop_box", None)
        if crop_box is None:
            return {"ok": False, "reason": "center_crop_box missing"}
    crop_box = tuple(crop_box)

    # 2) crops
    img0 = load_bgr(orig_path)
    img1 = load_bgr(aug_path)
    g0 = crop_gray(img0, crop_box)
    g1 = crop_gray(img1, crop_box)
    h, w = g0.shape[:2]
    crop_wh = (w, h)

    # 3) match
    t2 = time.time()
    nkp0 = nkp1 = None
    loftr_mean_conf = None

    if matcher == "loftr":
        if loftr is None:
            return {"ok": False, "reason": "matcher=loftr but LoFTR not available"}
        pts0, pts1, nmatch, s0, s1, loftr_mean_conf = loftr.match(g0, g1, max_side=loftr_max_side, conf_th=loftr_conf_th)
    else:
        pts0, pts1, nkp0, nkp1, nmatch = orb_match_points(g0, g1)
    t3 = time.time()

    # 4) H robust + sym error filter
    pts0_f, pts1_f, ninl, H, sym_mean = robust_h_inliers(
        pts0, pts1, thr_px=h_thr_px, use_usac_magsac=use_usac_magsac,
        sym_err_th=sym_err_th, sym_err_keep=sym_err_keep
    )
    t4 = time.time()

    if ninl < min_inliers:
        return {
            "ok": False,
            "reason": f"too few inliers after sym-filter: {ninl} (<{min_inliers})",
            "nmatch": int(nmatch),
            "ninliers": int(ninl),
            "loftr_mean_conf": loftr_mean_conf,
            "sym_mean": sym_mean,
            "t_moirepose": float(t1 - t0),
            "t_match": float(t3 - t2),
            "t_h": float(t4 - t3),
            "t_total": float(time.time() - t_total0),
        }

    # subsample inliers
    if max_refine_points > 0 and pts0_f.shape[0] > max_refine_points:
        idx = np.random.choice(pts0_f.shape[0], max_refine_points, replace=False)
        pts0_f = pts0_f[idx]
        pts1_f = pts1_f[idx]
        ninl = int(pts0_f.shape[0])

    # 5) a0 for legacy fronto-parallel lifting.  Prefer the value produced by
    # MP.calibrate_single because it now follows the crop-dependent Eq. (6).
    distances = aug_res.get("distances_by_poi", None)
    if not isinstance(distances, dict) or len(distances) == 0:
        return {"ok": False, "reason": "aug_res.distances_by_poi missing/empty"}

    d0 = float(np.mean(list(distances.values())))
    cfa_pitch = float(params["cfa_pitch_m"])
    f_cam = float(params["f_m"])
    half = max((min(h, w) / 2.0) - 4.0, 1.0)
    a0 = float(d0 * cfa_pitch / (f_cam + 1e-12) * half)
    try:
        a0 = float(aug_res.get("diagnostics", {}).get("a_m", a0))
    except Exception:
        pass

    # 6) K
    Kmat = crop_intrinsics(crop_wh, crop_box, use_full_k, full_fx_fy_cx_cy, params)

    # initial pose (from MP aug pose)
    pose_aug = aug_res.get("pose", None)
    if pose_aug is None:
        return {"ok": False, "reason": "aug_res.pose missing"}
    rvec0, tvec0 = mp_pose_to_opencv_rt(pose_aug)

    pose_orig = orig_res.get("pose", None)
    if pose_orig is None:
        return {"ok": False, "reason": "orig_res.pose missing"}
    rvec_orig, tvec_orig = mp_pose_to_opencv_rt(pose_orig)

    # RMS before (with a0)
    lift_valid_count = None
    if object_lift == "ray_plane":
        obj0, lift_valid = ray_plane_lift_crop_pixels_np(
            pts0_f,
            K_crop=Kmat,
            rvec_screen_to_cam=rvec_orig,
            tvec_screen_to_cam=tvec_orig,
        )
        if not np.all(lift_valid):
            pts0_f = pts0_f[lift_valid]
            pts1_f = pts1_f[lift_valid]
            obj0 = obj0[lift_valid]
            ninl = int(pts0_f.shape[0])
        lift_valid_count = int(ninl)
        if ninl < min_inliers:
            return {
                "ok": False,
                "reason": f"too few valid ray-plane lifted points: {ninl} (<{min_inliers})",
                "nmatch": int(nmatch), "ninliers": int(ninl),
                "t_moirepose": float(t1 - t0),
                "t_match": float(t3 - t2),
                "t_h": float(t4 - t3),
                "t_total": float(time.time() - t_total0),
            }
    else:
        obj0 = crop_px_to_screen_xy_np(pts0_f, crop_wh, a0)
    pred0 = project_points_cv(obj0, rvec0, tvec0, Kmat)
    rms_before = rms_px(pts1_f, pred0)

    # gate: if too bad, skip refine (prevents blow-ups)
    if gate_rms_before > 0 and rms_before > gate_rms_before:
        return {
            "ok": False,
            "reason": f"gated_by_rms_before: {rms_before:.2f} > {gate_rms_before}",
            "rms_before_px": float(rms_before),
            "nmatch": int(nmatch), "ninliers": int(ninl),
            "loftr_mean_conf": loftr_mean_conf,
            "sym_mean": sym_mean,
            "object_lift": object_lift,
            "t_moirepose": float(t1 - t0),
            "t_match": float(t3 - t2),
            "t_h": float(t4 - t3),
            "t_total": float(time.time() - t_total0),
        }

    # 7) refine
    t5 = time.time()
    torch_loss = None
    a_ref = a0

    if refine_backend == "torch":
        rvec_ref, tvec_ref, a_ref, torch_loss = refine_pose_torch(
            pts0_px_np=pts0_f,
            pts1_px_np=pts1_f,
            K_np=Kmat,
            rvec0_np=rvec0,
            tvec0_np=tvec0,
            a0=a0,
            crop_wh=crop_wh,
            optimize_scale_a=optimize_scale_a,
            device=refine_device,
            iters=torch_iters,
            huber_delta=huber_delta,
            method=torch_method,
            obj_pts_fixed_np=obj0 if object_lift == "ray_plane" else None,
        )
    else:
        # OpenCV pose-only refine
        rvec_ref, tvec_ref = cv2.solvePnPRefineLM(
            objectPoints=obj0,
            imagePoints=pts1_f.astype(np.float64),
            cameraMatrix=Kmat,
            distCoeffs=None,
            rvec=rvec0,
            tvec=tvec0
        )

    t6 = time.time()

    # RMS after.  If ray-plane lifting is selected, object points are fixed metric
    # screen points; if legacy fronto-parallel is selected, scale a may be refined.
    obj1 = obj0 if object_lift == "ray_plane" else crop_px_to_screen_xy_np(pts0_f, crop_wh, a_ref)
    pred1 = project_points_cv(obj1, rvec_ref, tvec_ref, Kmat)
    rms_after = rms_px(pts1_f, pred1)

    # pose sanity gate (prevents bad tails)
    C, u, v, w_axis = opencv_rt_to_mp_pose(rvec_ref, tvec_ref)
    if gate_abs_xyz_m > 0:
        if (abs(C[0]) > gate_abs_xyz_m) or (abs(C[1]) > gate_abs_xyz_m) or (abs(C[2]) > gate_abs_xyz_m):
            return {
                "ok": False,
                "reason": f"gated_by_pose: |C| > {gate_abs_xyz_m}m (C={C.tolist()})",
                "rms_before_px": float(rms_before),
                "rms_after_px": float(rms_after),
                "nmatch": int(nmatch), "ninliers": int(ninl),
                "t_moirepose": float(t1 - t0),
                "t_match": float(t3 - t2),
                "t_h": float(t4 - t3),
                "t_refine": float(t6 - t5),
                "t_total": float(time.time() - t_total0),
            }

    roll_keep = None
    try:
        roll_keep = float(pose_aug.get("roll_theta_c", None))
    except Exception:
        roll_keep = None

    refined_pose = {
        "x": float(C[0]), "y": float(C[1]), "z": float(C[2]),
        "u": (float(u[0]), float(u[1]), float(u[2])),
        "v": (float(v[0]), float(v[1]), float(v[2])),
        "w": (float(w_axis[0]), float(w_axis[1]), float(w_axis[2])),
        "roll_theta_c": roll_keep,
    }

    return {
        "ok": True,
        "pair": pair,
        "crop_box": crop_box,

        "nmatch": int(nmatch),
        "ninliers": int(ninl),
        "loftr_mean_conf": loftr_mean_conf,
        "sym_mean": sym_mean,

        "a0_m": float(a0),
        "a_ref_m": float(a_ref),
        "object_lift": object_lift,
        "ray_plane_valid_points": lift_valid_count,

        "rms_before_px": float(rms_before),
        "rms_after_px": float(rms_after),

        "refined_pose": refined_pose,
        "torch_final_loss": torch_loss,

        "t_moirepose": float(t1 - t0),
        "t_match": float(t3 - t2),
        "t_h": float(t4 - t3),
        "t_refine": float(t6 - t5),
        "t_total": float(time.time() - t_total0),
    }


# ============================================================
# Batch
# ============================================================
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--aug_dir", default="augmented_9x")
    ap.add_argument("--out_xlsx", default="hybrid_best.xlsx")
    ap.add_argument("--patch_mode", action="store_true")
    ap.add_argument("--patch_grid", default="5,5", help="patch grid rows,cols; e.g. 3,3 or 5,5")
    ap.add_argument("--patch_score_mode", default="reliability", choices=["reliability", "internal_rms"],
                    help="reliability = internal RMS + spectral confidence + physical/radial penalties")

    # matching
    ap.add_argument("--matcher", default="loftr", choices=["loftr", "orb"])
    ap.add_argument("--device", default="cuda")  # LoFTR device
    ap.add_argument("--loftr_pretrained", default="outdoor", choices=["outdoor", "indoor"])
    ap.add_argument("--loftr_max_side", type=int, default=512)
    ap.add_argument("--loftr_conf_th", type=float, default=0.3)

    # homography / inliers
    ap.add_argument("--use_usac_magsac", action="store_true")
    ap.add_argument("--h_thr_px", type=float, default=3.0)
    ap.add_argument("--sym_err_th", type=float, default=6.0)     # symmetric error threshold (px), 0 disables
    ap.add_argument("--sym_err_keep", type=int, default=400)     # keep top-K smallest sym errors, 0 disables
    ap.add_argument("--min_inliers", type=int, default=120)
    ap.add_argument("--max_refine_points", type=int, default=200)

    # refine
    ap.add_argument("--refine_backend", default="torch", choices=["torch", "opencv"])
    ap.add_argument("--refine_device", default="cuda")
    ap.add_argument("--torch_iters", type=int, default=10)
    ap.add_argument("--torch_method", default="adam", choices=["adam", "lbfgs"])
    ap.add_argument("--huber_delta", type=float, default=3.0)
    ap.add_argument("--optimize_scale_a", action="store_true")
    ap.add_argument("--object_lift", default="ray_plane", choices=["ray_plane", "fronto_parallel"],
                    help="ray_plane replaces the old Eq.(13) fronto-parallel object-point lift")

    # intrinsics
    ap.add_argument("--use_full_k", action="store_true")
    ap.add_argument("--K_full_fx_fy_cx_cy", nargs=4, type=float, default=None)

    # gating (tail-cut)
    ap.add_argument("--gate_rms_before", type=float, default=60.0)  # 0 disables
    ap.add_argument("--gate_abs_xyz_m", type=float, default=5.0)    # 0 disables

    args = ap.parse_args()

    if args.matcher == "loftr" and not _HAS_LOFTR:
        raise RuntimeError("LoFTR requested but kornia/torch not installed. Install kornia + torch(cuda).")
    if args.refine_backend == "torch" and not _HAS_TORCH:
        raise RuntimeError("Torch refine requested but torch not installed.")
    if args.device.startswith("cuda") and args.matcher == "loftr" and (not torch.cuda.is_available()):
        raise RuntimeError("LoFTR device=cuda requested but torch.cuda.is_available() is False.")
    if args.refine_backend == "torch" and args.refine_device.startswith("cuda") and (not torch.cuda.is_available()):
        raise RuntimeError("Torch refine_device=cuda requested but torch.cuda.is_available() is False.")

    full_fx_fy_cx_cy = None
    if args.use_full_k:
        if args.K_full_fx_fy_cx_cy is None:
            raise RuntimeError("--use_full_k requires --K_full_fx_fy_cx_cy fx fy cx cy")
        full_fx_fy_cx_cy = tuple(map(float, args.K_full_fx_fy_cx_cy))

    loftr = None
    if args.matcher == "loftr":
        loftr = LoFTRMatcher(device=args.device, pretrained=args.loftr_pretrained)

    root = Path(args.root)
    aug_dir = root / args.aug_dir
    out_xlsx = Path(args.out_xlsx)
    if not out_xlsx.is_absolute():
        out_xlsx = root / out_xlsx

    originals = find_originals(root)
    augmented = find_augmented(aug_dir)

    print(f"[INFO] root={root}")
    print(f"[INFO] originals={len(originals)} augmented_pairs={len(augmented)} patch_mode={args.patch_mode}")
    patch_grid = tuple(int(x.strip()) for x in args.patch_grid.split(","))
    if len(patch_grid) != 2:
        raise RuntimeError("--patch_grid must be rows,cols")
    print(f"[INFO] patch_grid={patch_grid} patch_score_mode={args.patch_score_mode}")
    print(f"[INFO] matcher={args.matcher} loftr_max_side={args.loftr_max_side} loftr_conf_th={args.loftr_conf_th}")
    print(f"[INFO] H: usac_magsac={args.use_usac_magsac} thr={args.h_thr_px} sym_th={args.sym_err_th} sym_keep={args.sym_err_keep}")
    print(f"[INFO] refine={args.refine_backend} optimize_scale_a={args.optimize_scale_a} iters={args.torch_iters} method={args.torch_method}")
    print(f"[INFO] object_lift={args.object_lift}")
    print(f"[INFO] use_full_k={args.use_full_k} K_full={full_fx_fy_cx_cy}")
    print(f"[INFO] gates: rms_before<{args.gate_rms_before}, |C|<{args.gate_abs_xyz_m}m")

    rows: List[Dict] = []

    for (idx, base, augtype, k, aug_path) in tqdm(augmented, desc="Hybrid BEST", unit="pair"):
        orig_path = originals.get((idx, base))
        if orig_path is None:
            rows.append({
                "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                "orig_path": None, "aug_path": str(aug_path),
                "error": f"missing original {idx}_{base}.jpg"
            })
            continue

        try:
            out = run_one_pair(
                orig_path=orig_path,
                aug_path=aug_path,
                params=MP.DEVICE_PARAMS,
                patch_mode=args.patch_mode,
                matcher=args.matcher,
                loftr=loftr,
                loftr_max_side=args.loftr_max_side,
                loftr_conf_th=args.loftr_conf_th,
                use_usac_magsac=args.use_usac_magsac,
                h_thr_px=args.h_thr_px,
                sym_err_th=args.sym_err_th,
                sym_err_keep=args.sym_err_keep,
                min_inliers=args.min_inliers,
                max_refine_points=args.max_refine_points,
                refine_backend=args.refine_backend,
                refine_device=args.refine_device,
                torch_iters=args.torch_iters,
                torch_method=args.torch_method,
                huber_delta=args.huber_delta,
                optimize_scale_a=args.optimize_scale_a,
                use_full_k=args.use_full_k,
                full_fx_fy_cx_cy=full_fx_fy_cx_cy,
                gate_rms_before=args.gate_rms_before,
                gate_abs_xyz_m=args.gate_abs_xyz_m,
                object_lift=args.object_lift,
                patch_grid=(patch_grid[0], patch_grid[1]),
                patch_score_mode=args.patch_score_mode,
            )

            if not out.get("ok", False):
                rows.append({
                    "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                    "orig_path": str(orig_path), "aug_path": str(aug_path),
                    "error": out.get("reason"),
                    "nmatch": out.get("nmatch"),
                    "ninliers": out.get("ninliers"),
                    "loftr_mean_conf": out.get("loftr_mean_conf"),
                    "sym_mean": out.get("sym_mean"),
                    "rms_before_px": out.get("rms_before_px"),
                    "rms_after_px": out.get("rms_after_px"),
                    "object_lift": out.get("object_lift"),
                    "t_moirepose": out.get("t_moirepose"),
                    "t_match": out.get("t_match"),
                    "t_h": out.get("t_h"),
                    "t_refine": out.get("t_refine"),
                    "t_total": out.get("t_total"),
                })
                continue

            pair = out["pair"]
            orig_res = pair["pair_result"]["orig"]
            aug_res = pair["pair_result"]["aug"]
            rp = out["refined_pose"]

            rows.append({
                "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                "orig_path": str(orig_path), "aug_path": str(aug_path),

                "nmatch": out.get("nmatch"),
                "ninliers": out.get("ninliers"),
                "loftr_mean_conf": out.get("loftr_mean_conf"),
                "sym_mean": out.get("sym_mean"),

                "a0_m": out.get("a0_m"),
                "a_ref_m": out.get("a_ref_m"),
                "object_lift": out.get("object_lift"),
                "ray_plane_valid_points": out.get("ray_plane_valid_points"),

                "hybrid_rms_before_px": out.get("rms_before_px"),
                "hybrid_rms_after_px": out.get("rms_after_px"),

                "ref_aug_x_m": rp.get("x"),
                "ref_aug_y_m": rp.get("y"),
                "ref_aug_z_m": rp.get("z"),
                "ref_aug_roll_rad_kept": rp.get("roll_theta_c"),

                "mp_aug_internal_rms_px": aug_res.get("rms_reproj_px", None),
                "torch_final_loss": out.get("torch_final_loss"),

                "t_moirepose": out.get("t_moirepose"),
                "t_match": out.get("t_match"),
                "t_h": out.get("t_h"),
                "t_refine": out.get("t_refine"),
                "t_total": out.get("t_total"),
            })

        except Exception as e:
            rows.append({
                "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                "orig_path": str(orig_path), "aug_path": str(aug_path),
                "error": f"{type(e).__name__}: {e}"
            })

    wb = Workbook()
    ws = wb.active
    ws.title = "hybrid_best"
    headers = sorted({k for r in rows for k in r.keys()})
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, None) for h in headers])
    autosize_columns(ws)

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_xlsx))
    print(f"\n[DONE] Saved: {out_xlsx}  rows={len(rows)}")


if __name__ == "__main__":
    main()
