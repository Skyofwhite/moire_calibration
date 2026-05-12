import argparse
import math
import re
import time
from pathlib import Path
from typing import Dict, List, Tuple

import cv2
import numpy as np
from tqdm import tqdm
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# ============================================================
# Filename patterns
# ============================================================
PAIR_RE = re.compile(
    r"^(?P<idx>\d{4})_(?P<base>gt|moire)_(?P<augtype>blackbox|rotation|translation)_(?P<k>\d+)\.jpg$",
    re.IGNORECASE
)
ORIG_RE = re.compile(r"^(?P<idx>\d{4})_(?P<base>gt|moire)\.jpg$", re.IGNORECASE)

# ============================================================
# File discovery
# ============================================================
def find_originals(test_dir: Path) -> Dict[Tuple[str, str], Path]:
    mapping: Dict[Tuple[str, str], Path] = {}
    for p in test_dir.iterdir():
        if p.is_file():
            m = ORIG_RE.match(p.name)
            if m:
                mapping[(m.group("idx"), m.group("base").lower())] = p
    return mapping


def find_augmented(aug_dir: Path) -> List[Tuple[str, str, str, str, Path]]:
    items = []
    for p in aug_dir.iterdir():
        if p.is_file():
            m = PAIR_RE.match(p.name)
            if m:
                items.append((m.group("idx"), m.group("base").lower(),
                              m.group("augtype").lower(), m.group("k"), p))
    items.sort(key=lambda t: (t[0], t[1], t[2], int(t[3])))
    return items

# ============================================================
# Utilities
# ============================================================
def autosize_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


def read_gray(path: Path) -> np.ndarray:
    img = cv2.imread(str(path), cv2.IMREAD_GRAYSCALE)
    if img is None:
        raise FileNotFoundError(str(path))
    return img


def resize_long_side(img: np.ndarray, max_side: int) -> Tuple[np.ndarray, float]:
    """
    Resize so max(H,W)=max_side, keep aspect ratio.
    Returns resized_img and scale s where resized = original*s.
    """
    h, w = img.shape[:2]
    s = max_side / max(h, w)
    if s >= 1.0:
        return img, 1.0
    out = cv2.resize(img, (int(round(w*s)), int(round(h*s))), interpolation=cv2.INTER_AREA)
    return out, s


def ensure_K_scaled(K: np.ndarray, scale: float) -> np.ndarray:
    Ks = K.copy().astype(np.float64)
    Ks[0, 0] *= scale
    Ks[1, 1] *= scale
    Ks[0, 2] *= scale
    Ks[1, 2] *= scale
    return Ks


def plane_points_from_ref_pixels(pts_ref_px: np.ndarray, sx_m_per_px: float, sy_m_per_px: float) -> np.ndarray:
    X = pts_ref_px[:, 0] * sx_m_per_px
    Y = pts_ref_px[:, 1] * sy_m_per_px
    Z = np.zeros_like(X)
    return np.stack([X, Y, Z], axis=1).astype(np.float64)


def reprojection_rms_px(object_points: np.ndarray, image_points: np.ndarray,
                        R: np.ndarray, t: np.ndarray,
                        K: np.ndarray, dist: np.ndarray) -> float:
    rvec, _ = cv2.Rodrigues(R)
    tvec = t.reshape(3, 1).astype(np.float64)
    proj, _ = cv2.projectPoints(object_points, rvec, tvec, K, dist)
    proj = proj.reshape(-1, 2)
    err = proj - image_points
    d2 = np.sum(err * err, axis=1)
    return float(np.sqrt(np.mean(d2)))


def positive_depth_ratio(object_points: np.ndarray, R: np.ndarray, t: np.ndarray) -> float:
    """
    For plane points (X,Y,0), camera depth is Zc = (R*[X,Y,0] + t)_z.
    """
    Xc = (R @ object_points.T).T + t.reshape(1, 3)
    zc = Xc[:, 2]
    return float(np.mean(zc > 0.0))


def rotation_angle_deg(R: np.ndarray) -> float:
    tr = float(np.trace(R))
    c = (tr - 1.0) / 2.0
    c = max(-1.0, min(1.0, c))
    return float(math.degrees(math.acos(c)))


def roll_from_R(R: np.ndarray) -> float:
    # same convention as earlier: roll = atan2(R[2,1], R[2,2])
    return float(math.atan2(R[2, 1], R[2, 2]))

# ============================================================
# Feature detection & matching
# ============================================================
def create_detector(detector: str):
    detector = detector.lower()
    if detector == "sift":
        if hasattr(cv2, "SIFT_create"):
            return cv2.SIFT_create()
        raise RuntimeError("SIFT not available. Install opencv-contrib-python or use ORB.")
    if detector == "orb":
        return cv2.ORB_create(nfeatures=5000)
    raise ValueError("detector must be 'orb' or 'sift'")


def match_features(detector_name: str, img0: np.ndarray, img1: np.ndarray):
    det = create_detector(detector_name)
    kp0, des0 = det.detectAndCompute(img0, None)
    kp1, des1 = det.detectAndCompute(img1, None)

    if des0 is None or des1 is None or len(kp0) < 20 or len(kp1) < 20:
        return np.zeros((0, 2), np.float64), np.zeros((0, 2), np.float64), 0, len(kp0), len(kp1)

    if detector_name == "sift":
        FLANN_INDEX_KDTREE = 1
        index_params = dict(algorithm=FLANN_INDEX_KDTREE, trees=5)
        search_params = dict(checks=50)
        flann = cv2.FlannBasedMatcher(index_params, search_params)
        knn = flann.knnMatch(des0, des1, k=2)
        matches = []
        for m, n in knn:
            if m.distance < 0.75 * n.distance:
                matches.append(m)
    else:
        bf = cv2.BFMatcher(cv2.NORM_HAMMING, crossCheck=True)
        matches = bf.match(des0, des1)
        matches = sorted(matches, key=lambda m: m.distance)[:2000]

    pts0 = np.float64([kp0[m.queryIdx].pt for m in matches]) if matches else np.zeros((0, 2), np.float64)
    pts1 = np.float64([kp1[m.trainIdx].pt for m in matches]) if matches else np.zeros((0, 2), np.float64)
    return pts0, pts1, len(matches), len(kp0), len(kp1)

# ============================================================
# Homography -> Pose via decomposition
# ============================================================
def select_pose_from_homography(
    H: np.ndarray,
    K: np.ndarray,
    dist: np.ndarray,
    pts_ref_px: np.ndarray,
    pts_img_px: np.ndarray,
    sx_m_per_px: float,
    sy_m_per_px: float,
    min_pos_depth_ratio: float = 0.7
):
    """
    Decompose homography and choose the best (R,t) based on:
      - positive depth ratio
      - reprojection RMS on same POI set
    """
    num, Rs, ts, ns = cv2.decomposeHomographyMat(H, K)
    if num <= 0:
        raise RuntimeError("decomposeHomographyMat failed")

    obj_pts = plane_points_from_ref_pixels(pts_ref_px, sx_m_per_px, sy_m_per_px)

    best = None
    for i in range(num):
        R = Rs[i].astype(np.float64)
        t = ts[i].reshape(3).astype(np.float64)

        # depth check on plane points
        pdr = positive_depth_ratio(obj_pts, R, t)

        # reprojection rms
        rms = reprojection_rms_px(obj_pts, pts_img_px, R, t, K, dist)

        cand = (rms, -pdr, i, R, t, pdr)  # sort: low rms, high pdr
        # keep all candidates; we'll pick best with pdr constraint if possible
        if best is None or cand < best:
            best = cand

    # try to pick best among candidates with sufficient positive depth ratio
    candidates = []
    for i in range(num):
        R = Rs[i].astype(np.float64)
        t = ts[i].reshape(3).astype(np.float64)
        pdr = positive_depth_ratio(obj_pts, R, t)
        rms = reprojection_rms_px(obj_pts, pts_img_px, R, t, K, dist)
        candidates.append((rms, -pdr, i, R, t, pdr))

    candidates.sort()
    # first candidate meeting depth constraint
    for rms, neg_pdr, i, R, t, pdr in candidates:
        if pdr >= min_pos_depth_ratio:
            return R, t, rms, pdr, i, num

    # fallback: best overall
    rms, neg_pdr, i, R, t, pdr = candidates[0]
    return R, t, rms, pdr, i, num

# ============================================================
# Per-pair processing
# ============================================================
def process_pair_planar_homography(
    orig_path: Path,
    aug_path: Path,
    detector: str,
    K: np.ndarray,
    dist: np.ndarray,
    sx_m_per_px: float,
    sy_m_per_px: float,
    max_side: int,
    h_ransac_thresh_px: float,
    min_pos_depth_ratio: float,
) -> Dict:
    g0 = read_gray(orig_path)
    g1 = read_gray(aug_path)

    g0s, s0 = resize_long_side(g0, max_side=max_side)
    g1s, s1 = resize_long_side(g1, max_side=max_side)

    # use same scale for both
    s = min(s0, s1)
    if s != s0:
        g0s = cv2.resize(g0, (int(round(g0.shape[1]*s)), int(round(g0.shape[0]*s))), interpolation=cv2.INTER_AREA)
    if s != s1:
        g1s = cv2.resize(g1, (int(round(g1.shape[1]*s)), int(round(g1.shape[0]*s))), interpolation=cv2.INTER_AREA)

    Ks = ensure_K_scaled(K, s)

    pts0, pts1, num_matches, n_kp0, n_kp1 = match_features(detector, g0s, g1s)
    if num_matches < 40:
        raise RuntimeError(f"Too few matches: {num_matches}")

    H, mask = cv2.findHomography(pts0, pts1, method=cv2.RANSAC, ransacReprojThreshold=h_ransac_thresh_px)
    if H is None or mask is None:
        raise RuntimeError("findHomography failed")

    mask = mask.ravel().astype(bool)
    pts0_f = pts0[mask]
    pts1_f = pts1[mask]
    if pts0_f.shape[0] < 30:
        raise RuntimeError(f"Too few homography inliers: {pts0_f.shape[0]}")

    # IMPORTANT: points are in resized coords. Convert screen scale accordingly.
    sx_scaled = (sx_m_per_px / s)
    sy_scaled = (sy_m_per_px / s)

    R_aug, t_aug, rms_aug, pdr, chosen_i, num_solutions = select_pose_from_homography(
        H=H.astype(np.float64),
        K=Ks,
        dist=dist,
        pts_ref_px=pts0_f,
        pts_img_px=pts1_f,
        sx_m_per_px=sx_scaled,
        sy_m_per_px=sy_scaled,
        min_pos_depth_ratio=min_pos_depth_ratio
    )

    roll_aug = roll_from_R(R_aug)
    dtheta_deg = rotation_angle_deg(R_aug)

    # anchor orig pose as identity for delta convenience
    dx, dy, dz = t_aug.tolist()

    return {
        "num_kp_orig": int(n_kp0),
        "num_kp_aug": int(n_kp1),
        "num_matches": int(num_matches),
        "num_h_inliers": int(pts0_f.shape[0]),
        "homography_num_solutions": int(num_solutions),
        "homography_chosen_solution": int(chosen_i),
        "pos_depth_ratio": float(pdr),

        "orig_rms_reproj_px": None,
        "aug_rms_reproj_px": float(rms_aug),

        "orig_x_m": 0.0, "orig_y_m": 0.0, "orig_z_m": 0.0,
        "orig_roll_rad": 0.0,

        "aug_x_m": float(t_aug[0]),
        "aug_y_m": float(t_aug[1]),
        "aug_z_m": float(t_aug[2]),
        "aug_roll_rad": float(roll_aug),

        "delta_x_m": float(dx),
        "delta_y_m": float(dy),
        "delta_z_m": float(dz),
        "delta_w_angle_deg": float(dtheta_deg),

        "max_side": int(max_side),
        "h_ransac_thresh_px": float(h_ransac_thresh_px),
        "min_pos_depth_ratio": float(min_pos_depth_ratio),
    }

# ============================================================
# Main
# ============================================================
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True)
    parser.add_argument("--aug_dir", default="augmented_9x")
    parser.add_argument("--out_xlsx", default="planar_homography_results_basic.xlsx")

    parser.add_argument("--detector", default="orb", choices=["orb", "sift"])
    parser.add_argument("--max_side", type=int, default=1280)

    # intrinsics
    parser.add_argument("--fx", type=float, required=True)
    parser.add_argument("--fy", type=float, required=True)
    parser.add_argument("--cx", type=float, required=True)
    parser.add_argument("--cy", type=float, required=True)

    # distortion
    parser.add_argument("--k1", type=float, default=0.0)
    parser.add_argument("--k2", type=float, default=0.0)
    parser.add_argument("--p1", type=float, default=0.0)
    parser.add_argument("--p2", type=float, default=0.0)
    parser.add_argument("--k3", type=float, default=0.0)

    # screen metric scale
    parser.add_argument("--sx_m_per_px", type=float, required=True)
    parser.add_argument("--sy_m_per_px", type=float, required=True)

    # homography ransac
    parser.add_argument("--h_ransac_thresh_px", type=float, default=3.0)
    parser.add_argument("--min_pos_depth_ratio", type=float, default=0.7)

    parser.add_argument("--skip_missing", action="store_true")
    args = parser.parse_args()

    Kmat = np.array([[args.fx, 0, args.cx],
                     [0, args.fy, args.cy],
                     [0, 0, 1]], dtype=np.float64)
    dist = np.array([args.k1, args.k2, args.p1, args.p2, args.k3], dtype=np.float64)

    root = Path(args.root)
    aug_dir = root / args.aug_dir
    out_xlsx = Path(args.out_xlsx)
    if not out_xlsx.is_absolute():
        out_xlsx = root / out_xlsx

    originals = find_originals(root)
    augmented = find_augmented(aug_dir)
    print(f"[INFO] Found originals: {len(originals)} | augmented pairs: {len(augmented)}")

    rows: List[Dict] = []
    for (idx, base, augtype, k, aug_path) in tqdm(augmented, desc="Planar Homography Calibration", unit="pair"):
        orig_path = originals.get((idx, base))
        if orig_path is None:
            if args.skip_missing:
                continue
            rows.append({
                "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                "orig_path": None, "aug_path": str(aug_path),
                "error": f"missing original {idx}_{base}.jpg"
            })
            continue

        t0 = time.time()
        try:
            metrics = process_pair_planar_homography(
                orig_path=orig_path,
                aug_path=aug_path,
                detector=args.detector,
                K=Kmat,
                dist=dist,
                sx_m_per_px=args.sx_m_per_px,
                sy_m_per_px=args.sy_m_per_px,
                max_side=args.max_side,
                h_ransac_thresh_px=args.h_ransac_thresh_px,
                min_pos_depth_ratio=args.min_pos_depth_ratio,
            )
            elapsed = time.time() - t0
            rows.append({
                "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                "orig_path": str(orig_path), "aug_path": str(aug_path),
                "elapsed_sec": round(elapsed, 3),
                **metrics
            })
        except Exception as e:
            rows.append({
                "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                "orig_path": str(orig_path), "aug_path": str(aug_path),
                "elapsed_sec": round(time.time() - t0, 3),
                "error": f"{type(e).__name__}: {e}"
            })

    wb = Workbook()
    ws = wb.active
    ws.title = "planar_homography_results"

    headers = sorted({k for r in rows for k in r.keys()})
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, None) for h in headers])

    autosize_columns(ws)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_xlsx))

    print(f"\n[DONE] Saved: {out_xlsx}")
    print(f"Rows: {len(rows)}")


if __name__ == "__main__":
    main()
