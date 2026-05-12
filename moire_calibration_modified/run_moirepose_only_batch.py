# run_moirepose_only_batch.py
"""
MoiréPose-only batch runner with safe checkpointing.

Key features:
- Writes each processed pair as a row to a checkpoint CSV (so 10+ hour runs don't get lost).
- Supports --resume to skip already processed pairs based on the checkpoint CSV.
- Converts the checkpoint CSV to the requested XLSX at the end (or via --finalize_only).
- Serializes non-Excel-safe values (tuple/list/dict/numpy scalars) to JSON strings.

Typical usage:
  # center crop baseline (no patch search)
  python run_moirepose_only_batch.py --root "...\datasets\test" --out_xlsx moirepose_center.xlsx

  # patch search baseline (25 patches)
  python run_moirepose_only_batch.py --root "...\datasets\test" --patch_mode --out_xlsx moirepose_patch.xlsx

  # if interrupted/crashed, resume:
  python run_moirepose_only_batch.py --root "...\datasets\test" --patch_mode --out_xlsx moirepose_patch.xlsx --resume

  # only convert existing checkpoint CSV -> XLSX:
  python run_moirepose_only_batch.py --root "...\datasets\test" --out_xlsx moirepose_patch.xlsx --finalize_only
"""
import argparse
import csv
import json
import os
import re
import sys
import time
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Set

import numpy as np
from tqdm import tqdm
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import moirepose_calib as MP

PAIR_RE = re.compile(
    r"^(?P<idx>\d{4})_(?P<base>gt|moire)_(?P<augtype>blackbox|rotation|translation)_(?P<k>\d+)\.jpg$",
    re.IGNORECASE,
)
ORIG_RE = re.compile(r"^(?P<idx>\d{4})_(?P<base>gt|moire)\.jpg$", re.IGNORECASE)

# Stable schema (helps CSV streaming + resume)
FIELDNAMES = [
    "idx", "base", "augtype", "aug_k",
    "orig_path", "aug_path",
    "elapsed_sec",

    "orig_internal_rms_px", "aug_internal_rms_px",

    "orig_x_m", "orig_y_m", "orig_z_m", "orig_roll_rad",
    "aug_x_m", "aug_y_m", "aug_z_m", "aug_roll_rad",

    "delta_w_angle_deg",
    "delta_x_m", "delta_y_m", "delta_z_m",

    # patch-mode debug
    "best_patch_index",
    "best_patch_score_rms_px",
    "center_patch_score_rms_px",
    "delta_score_center_minus_best_px",
    "chosen_patch_box",

    # error
    "error",
]

FORCE_TEXT_COLS = {"idx", "base", "augtype", "orig_path", "aug_path", "chosen_patch_box", "error"}


def autosize_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 70)


def find_originals(test_dir: Path) -> Dict[Tuple[str, str], Path]:
    mapping: Dict[Tuple[str, str], Path] = {}
    for p in test_dir.iterdir():
        if p.is_file():
            m = ORIG_RE.match(p.name)
            if m:
                mapping[(m.group("idx"), m.group("base").lower())] = p
    return mapping


def find_augmented(aug_dir: Path) -> List[Tuple[str, str, str, str, Path]]:
    items: List[Tuple[str, str, str, str, Path]] = []
    if not aug_dir.exists():
        return items
    for p in aug_dir.iterdir():
        if p.is_file():
            m = PAIR_RE.match(p.name)
            if m:
                items.append((m.group("idx"), m.group("base").lower(),
                              m.group("augtype").lower(), m.group("k"), p))
    items.sort(key=lambda t: (t[0], t[1], t[2], int(t[3])))
    return items


def safe_get(d: dict, k: str, default=None):
    try:
        return d.get(k, default)
    except Exception:
        return default


def make_key(idx: str, base: str, augtype: str, k: int) -> str:
    return f"{idx}|{base}|{augtype}|{k}"


def json_safe(v):
    """Convert arbitrary objects to JSON/CSV safe representation."""
    if v is None:
        return None

    # numpy scalar -> python scalar
    if isinstance(v, (np.integer, np.floating, np.bool_)):
        return v.item()

    # pathlib.Path
    if isinstance(v, Path):
        return str(v)

    # tuple/list/ndarray -> JSON string
    if isinstance(v, (tuple, list, np.ndarray)):
        return json.dumps([json_safe(x) for x in list(v)], ensure_ascii=False)

    # dict -> JSON string
    if isinstance(v, dict):
        return json.dumps({str(k): json_safe(val) for k, val in v.items()}, ensure_ascii=False)

    return v


def normalize_row(row: dict) -> dict:
    """Ensure all fieldnames exist and values are CSV-safe."""
    out = {k: None for k in FIELDNAMES}
    for k, v in row.items():
        if k in out:
            out[k] = json_safe(v)
        else:
            # ignore unknown keys to keep schema stable
            pass
    return out


def ensure_checkpoint_writer(csv_path: Path, fieldnames: List[str]) -> Tuple[csv.DictWriter, object, bool]:
    """
    Returns (writer, file_handle, is_new_file).
    """
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    is_new = not csv_path.exists() or csv_path.stat().st_size == 0
    f = open(csv_path, "a", newline="", encoding="utf-8")
    writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
    if is_new:
        writer.writeheader()
        f.flush()
        os.fsync(f.fileno())
    return writer, f, is_new


def load_processed_keys(csv_path: Path) -> Set[str]:
    if not csv_path.exists() or csv_path.stat().st_size == 0:
        return set()
    processed = set()
    with open(csv_path, "r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            try:
                idx = (r.get("idx") or "").strip()
                base = (r.get("base") or "").strip().lower()
                augtype = (r.get("augtype") or "").strip().lower()
                k = r.get("aug_k")
                if idx and base and augtype and k is not None and str(k).strip() != "":
                    processed.add(make_key(idx, base, augtype, int(float(k))))
            except Exception:
                # tolerate partially written/garbled lines
                continue
    return processed


def parse_cell_value(col: str, s: Optional[str]):
    if s is None:
        return None
    s = str(s).strip()
    if s == "" or s.lower() == "nan":
        return None
    if col in FORCE_TEXT_COLS:
        return s
    # keep JSON strings as text
    if (s.startswith("[") and s.endswith("]")) or (s.startswith("{") and s.endswith("}")):
        return s
    # numeric inference
    try:
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        return float(s)
    except Exception:
        return s


def finalize_csv_to_xlsx(csv_path: Path, out_xlsx: Path, sheet_name: str = "moirepose_only") -> None:
    if not csv_path.exists() or csv_path.stat().st_size == 0:
        raise RuntimeError(f"Checkpoint CSV not found or empty: {csv_path}")

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(FIELDNAMES)

    with open(csv_path, "r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            ws.append([parse_cell_value(col, r.get(col)) for col in FIELDNAMES])

    autosize_columns(ws)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_xlsx))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True, help="dataset root containing 0000_gt.jpg, 0000_moire.jpg, ...")
    ap.add_argument("--aug_dir", default="augmented_9x", help="subdir under root containing augmented pairs")
    ap.add_argument("--out_xlsx", default="moirepose_only.xlsx")
    ap.add_argument("--patch_mode", action="store_true", help="enable 5x5 patch search (min internal RMS)")

    # checkpointing / resume
    ap.add_argument("--checkpoint_csv", default=None, help="checkpoint CSV path (default: <out_xlsx>.checkpoint.csv)")
    ap.add_argument("--resume", action="store_true", help="resume from checkpoint CSV (skip already processed pairs)")
    ap.add_argument("--finalize_only", action="store_true", help="only convert checkpoint CSV -> XLSX; do not run inference")
    ap.add_argument("--flush_every", type=int, default=1, help="fsync checkpoint every N rows (default: 1)")

    args = ap.parse_args()

    root = Path(args.root)
    aug_dir = root / args.aug_dir
    out_xlsx = Path(args.out_xlsx)
    if not out_xlsx.is_absolute():
        out_xlsx = root / out_xlsx

    checkpoint_csv = Path(args.checkpoint_csv) if args.checkpoint_csv else out_xlsx.with_suffix(out_xlsx.suffix + ".checkpoint.csv")
    if not checkpoint_csv.is_absolute():
        checkpoint_csv = root / checkpoint_csv

    if args.finalize_only:
        print(f"[FINALIZE_ONLY] Converting checkpoint -> xlsx")
        print(f"[INFO] checkpoint_csv={checkpoint_csv}")
        print(f"[INFO] out_xlsx={out_xlsx}")
        finalize_csv_to_xlsx(checkpoint_csv, out_xlsx)
        print(f"[DONE] Saved: {out_xlsx}")
        return

    originals = find_originals(root)
    augmented = find_augmented(aug_dir)

    processed_keys: Set[str] = set()
    if args.resume:
        processed_keys = load_processed_keys(checkpoint_csv)
        print(f"[RESUME] loaded processed rows: {len(processed_keys)} from {checkpoint_csv}")

    # filter todo list
    todo = []
    for (idx, base, augtype, k, aug_path) in augmented:
        key = make_key(idx, base, augtype, int(k))
        if key not in processed_keys:
            todo.append((idx, base, augtype, k, aug_path))

    print(f"[INFO] root={root}")
    print(f"[INFO] aug_dir={aug_dir} total_pairs={len(augmented)} todo={len(todo)} patch_mode={args.patch_mode}")
    print(f"[INFO] checkpoint_csv={checkpoint_csv}")
    print(f"[INFO] out_xlsx={out_xlsx}")

    writer, f_handle, is_new = ensure_checkpoint_writer(checkpoint_csv, FIELDNAMES)
    rows_written_since_flush = 0

    center_patch_idx = 12  # 5x5 grid center (row-major)

    try:
        for (idx, base, augtype, k, aug_path) in tqdm(todo, desc="MoiréPose-only", unit="pair"):
            orig_path = originals.get((idx, base))
            if orig_path is None:
                row = {
                    "idx": idx, "base": base, "augtype": augtype, "aug_k": int(k),
                    "orig_path": None, "aug_path": str(aug_path),
                    "error": f"missing original {idx}_{base}.jpg",
                }
                writer.writerow(normalize_row(row))
                rows_written_since_flush += 1
                if rows_written_since_flush >= args.flush_every:
                    f_handle.flush()
                    os.fsync(f_handle.fileno())
                    rows_written_since_flush = 0
                continue

            t0 = time.time()
            try:
                out = MP.calibrate_pair(str(orig_path), str(aug_path), MP.DEVICE_PARAMS, patch_mode=args.patch_mode)
                t1 = time.time()

                pr = out.get("pair_result", {})
                o = pr.get("orig", {})
                a = pr.get("aug", {})

                row = {
                    "idx": idx, "base": base, "augtype": augtype, "aug_k": int(k),
                    "orig_path": str(orig_path), "aug_path": str(aug_path),
                    "elapsed_sec": float(t1 - t0),

                    "orig_internal_rms_px": safe_get(o, "rms_reproj_px"),
                    "aug_internal_rms_px": safe_get(a, "rms_reproj_px"),

                    "orig_x_m": safe_get(safe_get(o, "pose", {}), "x"),
                    "orig_y_m": safe_get(safe_get(o, "pose", {}), "y"),
                    "orig_z_m": safe_get(safe_get(o, "pose", {}), "z"),
                    "orig_roll_rad": safe_get(safe_get(o, "pose", {}), "roll_theta_c"),

                    "aug_x_m": safe_get(safe_get(a, "pose", {}), "x"),
                    "aug_y_m": safe_get(safe_get(a, "pose", {}), "y"),
                    "aug_z_m": safe_get(safe_get(a, "pose", {}), "z"),
                    "aug_roll_rad": safe_get(safe_get(a, "pose", {}), "roll_theta_c"),

                    "delta_w_angle_deg": safe_get(pr, "delta_w_angle_deg"),
                }
                dpos = safe_get(pr, "delta_position_m")
                if isinstance(dpos, (list, tuple)) and len(dpos) == 3:
                    row["delta_x_m"], row["delta_y_m"], row["delta_z_m"] = dpos

                # Patch debug summary (optional)
                if args.patch_mode:
                    dbg = out.get("debug", {})
                    pd = dbg.get("patch_debug", {})
                    row["best_patch_index"] = pd.get("best_patch_index")
                    row["best_patch_score_rms_px"] = pd.get("best_patch_score_rms_reproj_px")
                    row["chosen_patch_box"] = dbg.get("chosen_patch_box")

                    center_info = pd.get(center_patch_idx, None)
                    row["center_patch_score_rms_px"] = center_info.get("score_rms_reproj_px") if isinstance(center_info, dict) else None

                    if row.get("best_patch_score_rms_px") is not None and row.get("center_patch_score_rms_px") is not None:
                        row["delta_score_center_minus_best_px"] = (
                            float(row["center_patch_score_rms_px"]) - float(row["best_patch_score_rms_px"])
                        )

                writer.writerow(normalize_row(row))
                rows_written_since_flush += 1

            except Exception as e:
                t1 = time.time()
                row = {
                    "idx": idx, "base": base, "augtype": augtype, "aug_k": int(k),
                    "orig_path": str(orig_path), "aug_path": str(aug_path),
                    "elapsed_sec": float(t1 - t0),
                    "error": f"{type(e).__name__}: {e}",
                }
                writer.writerow(normalize_row(row))
                rows_written_since_flush += 1

            if rows_written_since_flush >= args.flush_every:
                f_handle.flush()
                os.fsync(f_handle.fileno())
                rows_written_since_flush = 0

    except KeyboardInterrupt:
        print("\n[WARN] Interrupted by user (Ctrl+C). Checkpoint CSV is preserved; you can resume with --resume.", file=sys.stderr)

    finally:
        try:
            f_handle.flush()
            os.fsync(f_handle.fileno())
        except Exception:
            pass
        try:
            f_handle.close()
        except Exception:
            pass

    # Finalize
    print(f"[FINALIZE] Converting checkpoint -> xlsx")
    print(f"[INFO] checkpoint_csv={checkpoint_csv}")
    print(f"[INFO] out_xlsx={out_xlsx}")
    finalize_csv_to_xlsx(checkpoint_csv, out_xlsx)
    print(f"[DONE] Saved: {out_xlsx}")
    print(f"[INFO] If you need to resume later, re-run with --resume (it will skip rows already in the checkpoint CSV).")


if __name__ == "__main__":
    main()
