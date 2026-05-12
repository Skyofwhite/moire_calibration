import argparse
import re
import time
from pathlib import Path
from typing import Dict, List, Tuple, Optional

import cv2
import numpy as np
from tqdm import tqdm
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# ---- import your MoiréPose implementation ----
import moirepose_calib as MP  # expects moirepose_calib.py in same folder

# Optional (GPU) matcher: LoFTR (Kornia)
try:
    import torch
    import kornia.feature as KF
    _HAS_LOFTR = True
except Exception:
    torch = None
    KF = None
    _HAS_LOFTR = False


PAIR_RE = re.compile(
    r"^(?P<idx>\d{4})_(?P<base>gt|moire)_(?P<augtype>blackbox|rotation|translation)_(?P<k>\d+)\.jpg$",
    re.IGNORECASE
)
ORIG_RE = re.compile(r"^(?P<idx>\d{4})_(?P<base>gt|moire)\.jpg$", re.IGNORECASE)


# -----------------------------
# Excel formatting helper
# -----------------------------
def autosize_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


# -----------------------------
# Pair discovery
# -----------------------------
def find_originals(test_dir: Path) -> Dict[Tuple[str, str], Path]:
    mapping: Dict[Tuple[str, str], Path] = {}
    for p in test_dir.iterdir():
        if p.is_file():
            m = ORIG_RE.match(p.name)
            if m:
                mapping[(m.group("idx"), m.group("base").lower())] = p
    return mapping


def find_augmented(aug_dir: Path) -> List[Tuple[str, str, str, str, Path]]:
    items = []
    if not aug_dir.exists():
        return items
    for p in aug_dir.iterdir():
        if p.is_file():
            m = PAIR_RE.match(p.name)
            if m:
                items.append((m.group("idx"), m.group("base").lower(),
                              m.group("augtype").lower(), m.group("k"), p))
    items.sort(key=lambda t: (t[0], t[1], t[2], int(t[3])))
    return items


# -----------------------------
# IO and crop
# -----------------------------
def load_bgr(path: Path) -> np.ndarray:
    img = cv2.imread(str(path), cv2.IMREAD_COLOR)
    if img is None:
        raise FileNotFoundError(str(path))
    return img


def crop_gray(img_bgr: np.ndarray, box: Tuple[int, int, int, int]) -> np.ndarray:
    x0, y0, x1, y1 = box
    crop = img_bgr[y0:y1, x0:x1]
    if crop.size == 0:
        raise ValueError(f"Empty crop for box={box} on image size={img_bgr.shape}")
    return cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)


# -----------------------------
# ORB matching (CPU)
# -----------------------------
def orb_match_points(
    g0: np.ndarray,
    g1: np.ndarray,
    nfeatures: int = 3000,
    max_matches: int = 1500
):
    orb = cv2.ORB_create(nfeatures=nfeatures)
    kp0, des0 = orb.detectAndCompute(g0, None)
    kp1, des1 = orb.detectAndCompute(g1, None)

    nkp0, nkp1 = len(kp0), len(kp1)
    if des0 is None or des1 is None or nkp0 < 30 or nkp1 < 30:
        return (np.zeros((0, 2), np.float64),
                np.zeros((0, 2), np.float64),
                nkp0, nkp1, 0)

    bf = cv2.BFMatcher(cv2.NORM_HAMMING, crossCheck=True)
    matches = bf.match(des0, des1)
    matches = sorted(matches, key=lambda m: m.distance)[:max_matches]

    pts0 = np.float64([kp0[m.queryIdx].pt for m in matches]) if matches else np.zeros((0, 2), np.float64)
    pts1 = np.float64([kp1[m.trainIdx].pt for m in matches]) if matches else np.zeros((0, 2), np.float64)
    return pts0, pts1, nkp0, nkp1, len(matches)


# -----------------------------
# LoFTR matching (GPU via Kornia)
# -----------------------------
class LoFTRMatcher:
    def __init__(self, device: str = "cuda", pretrained: str = "outdoor"):
        if not _HAS_LOFTR:
            raise RuntimeError("LoFTR requested but kornia/torch is not installed.")
        if device.startswith("cuda"):
            if not torch.cuda.is_available():
                raise RuntimeError("CUDA requested but torch.cuda.is_available() is False.")
        self.device = torch.device(device)
        self.matcher = KF.LoFTR(pretrained=pretrained).to(self.device).eval()

    @torch.inference_mode()
    def match(self, g0: np.ndarray, g1: np.ndarray, max_side: int = 640):
        """
        Returns (mkpts0, mkpts1, nmatch, s0, s1)
        mkpts are in ORIGINAL crop pixel coordinates.
        """
        def resize_keep(img, max_side_):
            h, w = img.shape[:2]
            s = max_side_ / max(h, w)
            if s >= 1.0:
                return img, 1.0
            out = cv2.resize(img, (int(round(w * s)), int(round(h * s))), interpolation=cv2.INTER_AREA)
            return out, s

        g0r, s0 = resize_keep(g0, max_side)
        g1r, s1 = resize_keep(g1, max_side)

        t0 = torch.from_numpy(g0r).float()[None, None] / 255.0
        t1 = torch.from_numpy(g1r).float()[None, None] / 255.0
        t0 = t0.to(self.device)
        t1 = t1.to(self.device)

        out = self.matcher({"image0": t0, "image1": t1})
        mkpts0 = out["keypoints0"].detach().cpu().numpy()
        mkpts1 = out["keypoints1"].detach().cpu().numpy()

        # scale back to original crop coords
        mkpts0[:, 0] /= s0
        mkpts0[:, 1] /= s0
        mkpts1[:, 0] /= s1
        mkpts1[:, 1] /= s1

        return mkpts0.astype(np.float64), mkpts1.astype(np.float64), int(mkpts0.shape[0]), float(s0), float(s1)


# -----------------------------
# RANSAC inliers via Homography (CPU)
# -----------------------------
def ransac_h_inliers(pts0: np.ndarray, pts1: np.ndarray, thr_px: float = 3.0):
    if pts0.shape[0] < 50:
        return np.zeros((0, 2), np.float64), np.zeros((0, 2), np.float64), 0
    H, mask = cv2.findHomography(pts0, pts1, cv2.RANSAC, thr_px)
    if H is None or mask is None:
        return np.zeros((0, 2), np.float64), np.zeros((0, 2), np.float64), 0
    m = mask.ravel().astype(bool)
    return pts0[m], pts1[m], int(m.sum())


# -----------------------------
# Mapping crop pixels -> MoiréPose screen plane (X,Y,0)
# -----------------------------
def crop_px_to_screen_xy(px: np.ndarray, crop_wh: Tuple[int, int], a_m: float, margin_px: int = 4):
    """
    Linear mapping:
      X = (x - cx) * s
      Y = -(y - cy) * s
      Z = 0
    where s = a / half_size.
    """
    w, h = crop_wh
    cx, cy = w / 2.0, h / 2.0

    half = (min(w, h) / 2.0) - float(margin_px)
    if half <= 1:
        half = min(w, h) / 2.0

    s = a_m / (half + 1e-12)
    x = (px[:, 0] - cx) * s
    y = -(px[:, 1] - cy) * s
    z = np.zeros_like(x)
    return np.stack([x, y, z], axis=1).astype(np.float64)


# -----------------------------
# Pose conversions: MP <-> OpenCV
# -----------------------------
def mp_pose_to_opencv_rt(pose: dict) -> Tuple[np.ndarray, np.ndarray]:
    """
    MP convention (from moirepose_calib.py):
      - pose has camera center C=(x,y,z) in "screen world"
      - axes u,v,w as columns of R_wc
      - world->camera: R_cw = R_wc^T
      - Xc = R_cw (Xw - C)
    OpenCV: Xc = R * Xw + t  =>  t = -R * C
    """
    C = np.array([pose["x"], pose["y"], pose["z"]], dtype=np.float64).reshape(3, 1)
    u = np.array(pose["u"], dtype=np.float64).reshape(3)
    v = np.array(pose["v"], dtype=np.float64).reshape(3)
    w = np.array(pose["w"], dtype=np.float64).reshape(3)

    R_wc = np.stack([u, v, w], axis=1)
    R_cw = R_wc.T
    t = -(R_cw @ C)
    rvec, _ = cv2.Rodrigues(R_cw)
    return rvec.astype(np.float64), t.astype(np.float64)


def opencv_rt_to_mp_pose(rvec: np.ndarray, tvec: np.ndarray):
    R_cw, _ = cv2.Rodrigues(rvec.reshape(3, 1))
    R_wc = R_cw.T
    C = -(R_wc @ tvec.reshape(3, 1))
    u = R_wc[:, 0]
    v = R_wc[:, 1]
    w = R_wc[:, 2]
    return C.reshape(3), u.reshape(3), v.reshape(3), w.reshape(3)


# -----------------------------
# Intrinsics consistent with MP
# -----------------------------
def crop_intrinsics_from_params(params: dict, crop_wh: Tuple[int, int]) -> np.ndarray:
    f_px = MP.focal_px(float(params["f_m"]), float(params["cfa_pitch_m"]))  # f_m / cfa_pitch
    w, h = crop_wh
    cx, cy = w / 2.0, h / 2.0
    return np.array([[f_px, 0.0, cx],
                     [0.0, f_px, cy],
                     [0.0, 0.0, 1.0]], dtype=np.float64)


def rms_px(obs: np.ndarray, pred: np.ndarray) -> float:
    e = obs - pred
    return float(np.sqrt(np.mean(np.sum(e * e, axis=1))))


def project_points_cv(obj_pts: np.ndarray, rvec: np.ndarray, tvec: np.ndarray, K: np.ndarray) -> np.ndarray:
    proj, _ = cv2.projectPoints(obj_pts, rvec, tvec, K, None)
    return proj.reshape(-1, 2)


# -----------------------------
# Core: hybrid refine one pair
# -----------------------------
def refine_aug_pose_hybrid(
    orig_path: Path,
    aug_path: Path,
    params: dict,
    patch_mode: bool,
    matcher: str,
    loftr: Optional[LoFTRMatcher],
    loftr_max_side: int,
    h_thr_px: float,
    min_inliers: int,
    max_refine_points: int,
    orb_nfeatures: int,
    orb_max_matches: int,
) -> Dict:
    t_total0 = time.time()

    # 1) MoiréPose (initial)
    t0 = time.time()
    pair = MP.calibrate_pair(str(orig_path), str(aug_path), params, patch_mode=patch_mode)
    t1 = time.time()

    # Validate outputs robustly
    if "pair_result" not in pair or "orig" not in pair["pair_result"] or "aug" not in pair["pair_result"]:
        return {"ok": False, "reason": "MP.calibrate_pair returned unexpected structure", "pair": pair}

    orig_res = pair["pair_result"]["orig"]
    aug_res = pair["pair_result"]["aug"]

    # crop box used by MP
    if patch_mode:
        crop_box = None
        if "debug" in pair and isinstance(pair["debug"], dict):
            crop_box = pair["debug"].get("chosen_patch_box", None)
        if crop_box is None:
            return {"ok": False, "reason": "patch_mode=True but debug.chosen_patch_box missing", "pair": pair}
    else:
        crop_box = orig_res.get("center_crop_box", None)
        if crop_box is None:
            return {"ok": False, "reason": "center_crop_box missing from MP result", "pair": pair}

    # 2) Load crops
    img0 = load_bgr(orig_path)
    img1 = load_bgr(aug_path)
    g0 = crop_gray(img0, tuple(crop_box))
    g1 = crop_gray(img1, tuple(crop_box))
    h, w = g0.shape[:2]
    crop_wh = (w, h)

    # 3) Matching (ORB or LoFTR)
    t2 = time.time()
    nkp0 = nkp1 = None
    loftr_s0 = loftr_s1 = None

    if matcher == "loftr":
        if loftr is None:
            return {"ok": False, "reason": "matcher=loftr but LoFTR is not available (install kornia/torch)", "pair": pair}
        pts0, pts1, nmatch, loftr_s0, loftr_s1 = loftr.match(g0, g1, max_side=loftr_max_side)
    else:
        pts0, pts1, nkp0, nkp1, nmatch = orb_match_points(
            g0, g1, nfeatures=orb_nfeatures, max_matches=orb_max_matches
        )
    t3 = time.time()

    # 4) Homography RANSAC inliers
    pts0_f, pts1_f, ninl = ransac_h_inliers(pts0, pts1, thr_px=h_thr_px)
    t4 = time.time()

    if ninl < min_inliers:
        return {
            "ok": False,
            "reason": f"too few inliers for refinement: {ninl} (<{min_inliers})",
            "pair": pair,
            "matcher": matcher,
            "nkp_orig": nkp0, "nkp_aug": nkp1,
            "nmatch": nmatch, "ninliers": ninl,
            "t_moirepose": float(t1 - t0),
            "t_match": float(t3 - t2),
            "t_h": float(t4 - t3),
            "t_total": float(time.time() - t_total0),
            "loftr_s0": loftr_s0, "loftr_s1": loftr_s1
        }

    # Subsample inliers for speed (LM cost)
    if max_refine_points > 0 and pts0_f.shape[0] > max_refine_points:
        idxs = np.random.choice(pts0_f.shape[0], max_refine_points, replace=False)
        pts0_f = pts0_f[idxs]
        pts1_f = pts1_f[idxs]
        ninl = int(pts0_f.shape[0])

    # 5) Compute MP scale 'a' consistent with MP internals
    if "distances_by_poi" not in aug_res:
        return {"ok": False, "reason": "aug_res.distances_by_poi missing", "pair": pair}

    distances = aug_res["distances_by_poi"]
    if not isinstance(distances, dict) or len(distances) == 0:
        return {"ok": False, "reason": "aug_res.distances_by_poi invalid/empty", "pair": pair}

    d0 = float(np.mean(list(distances.values())))
    cfa_pitch = float(params["cfa_pitch_m"])
    f_cam = float(params["f_m"])
    Lc = float(min(h, w) * cfa_pitch)
    a = float(d0 * Lc / (6.0 * f_cam + 1e-12))

    obj_pts = crop_px_to_screen_xy(pts0_f, crop_wh=crop_wh, a_m=a, margin_px=4)
    img_pts = pts1_f.astype(np.float64)

    # 6) Intrinsics and initial RT from MP
    Kmat = crop_intrinsics_from_params(params, crop_wh)
    if "pose" not in aug_res:
        return {"ok": False, "reason": "aug_res.pose missing", "pair": pair}

    rvec0, tvec0 = mp_pose_to_opencv_rt(aug_res["pose"])

    # RMS before on same point set
    pred0 = project_points_cv(obj_pts, rvec0, tvec0, Kmat)
    rms_before = rms_px(img_pts, pred0)

    # 7) LM refinement (CPU OpenCV)
    t5 = time.time()
    rvec_ref, tvec_ref = cv2.solvePnPRefineLM(
        objectPoints=obj_pts,
        imagePoints=img_pts,
        cameraMatrix=Kmat,
        distCoeffs=None,
        rvec=rvec0,
        tvec=tvec0
    )
    t6 = time.time()

    pred1 = project_points_cv(obj_pts, rvec_ref, tvec_ref, Kmat)
    rms_after = rms_px(img_pts, pred1)

    # Convert back to MP-style pose
    C, u, v, w_axis = opencv_rt_to_mp_pose(rvec_ref, tvec_ref)

    # roll is derived from moiré model; keep original for reporting
    roll_keep = None
    try:
        roll_keep = float(aug_res["pose"]["roll_theta_c"])
    except Exception:
        roll_keep = None

    refined_pose = {
        "x": float(C[0]), "y": float(C[1]), "z": float(C[2]),
        "u": (float(u[0]), float(u[1]), float(u[2])),
        "v": (float(v[0]), float(v[1]), float(v[2])),
        "w": (float(w_axis[0]), float(w_axis[1]), float(w_axis[2])),
        "roll_theta_c": roll_keep,
    }

    return {
        "ok": True,
        "pair": pair,

        "matcher": matcher,
        "loftr_max_side": loftr_max_side if matcher == "loftr" else None,
        "loftr_s0": loftr_s0, "loftr_s1": loftr_s1,
        "nkp_orig": nkp0, "nkp_aug": nkp1,
        "nmatch": int(nmatch),
        "ninliers": int(ninl),

        "a_m": float(a),

        "rms_before_px": float(rms_before),
        "rms_after_px": float(rms_after),

        "refined_pose": refined_pose,

        "t_moirepose": float(t1 - t0),
        "t_match": float(t3 - t2),
        "t_h": float(t4 - t3),
        "t_refine": float(t6 - t5),
        "t_total": float(time.time() - t_total0),
    }


# -----------------------------
# Main batch runner
# -----------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--aug_dir", default="augmented_9x")
    ap.add_argument("--out_xlsx", default="moirepose_hybrid_refine.xlsx")

    # MP mode
    ap.add_argument("--patch_mode", action="store_true")

    # matching
    ap.add_argument("--matcher", default="loftr", choices=["loftr", "orb"])
    ap.add_argument("--h_thr_px", type=float, default=3.0)
    ap.add_argument("--min_inliers", type=int, default=80)
    ap.add_argument("--max_refine_points", type=int, default=300)

    # ORB params
    ap.add_argument("--orb_nfeatures", type=int, default=3000)
    ap.add_argument("--orb_max_matches", type=int, default=1500)

    # LoFTR params
    ap.add_argument("--device", default="cuda")  # "cuda" or "cpu"
    ap.add_argument("--loftr_pretrained", default="outdoor", choices=["outdoor", "indoor"])
    ap.add_argument("--loftr_max_side", type=int, default=640)

    args = ap.parse_args()

    # Setup matcher if needed
    loftr = None
    if args.matcher == "loftr":
        if not _HAS_LOFTR:
            raise RuntimeError("matcher=loftr but kornia/torch not installed. Install kornia + torch(cuda).")
        loftr = LoFTRMatcher(device=args.device, pretrained=args.loftr_pretrained)

    root = Path(args.root)
    aug_dir = root / args.aug_dir
    out_xlsx = Path(args.out_xlsx)
    if not out_xlsx.is_absolute():
        out_xlsx = root / out_xlsx

    originals = find_originals(root)
    augmented = find_augmented(aug_dir)

    print(f"[INFO] root={root}")
    print(f"[INFO] originals={len(originals)} augmented_pairs={len(augmented)} patch_mode={args.patch_mode}")
    print(f"[INFO] matcher={args.matcher} device={args.device} loftr_max_side={args.loftr_max_side} max_refine_points={args.max_refine_points}")

    rows: List[Dict] = []
    for (idx, base, augtype, k, aug_path) in tqdm(augmented, desc="MoiréPose Hybrid (GPU match + LM refine)", unit="pair"):
        orig_path = originals.get((idx, base))
        if orig_path is None:
            rows.append({
                "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                "orig_path": None, "aug_path": str(aug_path),
                "error": f"missing original {idx}_{base}.jpg"
            })
            continue

        try:
            out = refine_aug_pose_hybrid(
                orig_path=orig_path,
                aug_path=aug_path,
                params=MP.DEVICE_PARAMS,
                patch_mode=args.patch_mode,
                matcher=args.matcher,
                loftr=loftr,
                loftr_max_side=args.loftr_max_side,
                h_thr_px=args.h_thr_px,
                min_inliers=args.min_inliers,
                max_refine_points=args.max_refine_points,
                orb_nfeatures=args.orb_nfeatures,
                orb_max_matches=args.orb_max_matches
            )

            if not out["ok"]:
                rows.append({
                    "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                    "orig_path": str(orig_path), "aug_path": str(aug_path),
                    "error": out.get("reason"),
                    "matcher": out.get("matcher"),
                    "nmatch": out.get("nmatch"),
                    "ninliers": out.get("ninliers"),
                    "t_moirepose": out.get("t_moirepose"),
                    "t_match": out.get("t_match"),
                    "t_h": out.get("t_h"),
                    "t_total": out.get("t_total"),
                    "loftr_s0": out.get("loftr_s0"),
                    "loftr_s1": out.get("loftr_s1"),
                })
                continue

            pair = out["pair"]
            orig_res = pair["pair_result"]["orig"]
            aug_res = pair["pair_result"]["aug"]
            rp = out["refined_pose"]

            # MP internal RMS (4 POIs)
            mp_rms_orig_internal = orig_res.get("rms_reproj_px", None)
            mp_rms_aug_internal = aug_res.get("rms_reproj_px", None)

            # MP pose
            mp_pose_aug = aug_res.get("pose", {})
            mp_aug_x = mp_pose_aug.get("x", None)
            mp_aug_y = mp_pose_aug.get("y", None)
            mp_aug_z = mp_pose_aug.get("z", None)
            mp_aug_roll = mp_pose_aug.get("roll_theta_c", None)

            # orig pose for delta (approx)
            mp_pose_orig = orig_res.get("pose", {})
            mp_orig_x = mp_pose_orig.get("x", None)
            mp_orig_y = mp_pose_orig.get("y", None)
            mp_orig_z = mp_pose_orig.get("z", None)

            rows.append({
                "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                "orig_path": str(orig_path), "aug_path": str(aug_path),

                "matcher": out["matcher"],
                "loftr_max_side": out.get("loftr_max_side"),
                "loftr_s0": out.get("loftr_s0"),
                "loftr_s1": out.get("loftr_s1"),

                "nkp_orig": out.get("nkp_orig"),
                "nkp_aug": out.get("nkp_aug"),
                "nmatch": out.get("nmatch"),
                "ninliers": out.get("ninliers"),
                "max_refine_points": args.max_refine_points,

                "a_m": out.get("a_m"),

                "mp_aug_x_m": mp_aug_x,
                "mp_aug_y_m": mp_aug_y,
                "mp_aug_z_m": mp_aug_z,
                "mp_aug_roll_rad": mp_aug_roll,

                "mp_orig_internal_rms_px": mp_rms_orig_internal,
                "mp_aug_internal_rms_px": mp_rms_aug_internal,

                "hybrid_rms_before_px": out.get("rms_before_px"),
                "hybrid_rms_after_px": out.get("rms_after_px"),

                "ref_aug_x_m": rp.get("x"),
                "ref_aug_y_m": rp.get("y"),
                "ref_aug_z_m": rp.get("z"),
                "ref_aug_roll_rad_kept": rp.get("roll_theta_c"),

                "ref_delta_x_m": (rp.get("x") - mp_orig_x) if (rp.get("x") is not None and mp_orig_x is not None) else None,
                "ref_delta_y_m": (rp.get("y") - mp_orig_y) if (rp.get("y") is not None and mp_orig_y is not None) else None,
                "ref_delta_z_m": (rp.get("z") - mp_orig_z) if (rp.get("z") is not None and mp_orig_z is not None) else None,

                "t_moirepose": out.get("t_moirepose"),
                "t_match": out.get("t_match"),
                "t_h": out.get("t_h"),
                "t_refine": out.get("t_refine"),
                "t_total": out.get("t_total"),
            })

        except Exception as e:
            rows.append({
                "idx": idx, "base": base, "augtype": augtype, "aug_k": k,
                "orig_path": str(orig_path), "aug_path": str(aug_path),
                "error": f"{type(e).__name__}: {e}"
            })

    # Write XLSX
    wb = Workbook()
    ws = wb.active
    ws.title = "moirepose_hybrid_gpu"

    headers = sorted({k for r in rows for k in r.keys()})
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, None) for h in headers])

    autosize_columns(ws)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_xlsx))

    print(f"\n[DONE] Saved: {out_xlsx}")
    print(f"Rows: {len(rows)}")


if __name__ == "__main__":
    main()
